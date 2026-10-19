"""Build DDC_Excel_Kit.xlsx: a workbook with a Config table and all Power Query (M) code
ready to paste. Excel cannot be scripted to create queries without Excel itself, so the
queries are listed in order on the 'Queries' sheet; docs/POWER_QUERY.md has the 2-minute steps.

    python build_kit.py DDC_Excel_Kit.xlsx
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
PBI = HERE.parent / "powerbi"
ORDER = [  # (query name, source file, purpose)
    ("ExportFolder", HERE / "ExportFolder.pq", "Reads the folder path from Config"),
    ("PreferredFormat", HERE / "PreferredFormat.pq", "Reads the preferred file format from Config"),
    ("RawElements", PBI / "RawElements.pq", "Stacks all export files of one format"),
    ("Elements", PBI / "Elements.pq", "Standard column names (ElementId, Category, ...)"),
    ("QuantitiesByType", PBI / "QuantitiesByType.pq", "Count and sums per category / type / level"),
    ("Completeness", PBI / "Completeness.pq", "Share of filled values per column"),
]


def build(target):
    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    steps = [
        "DDC Excel kit: read converter exports with Power Query (no Python, no account)",
        "",
        "1. Open the Config sheet and type the folder with your export (xlsx, csv or jsonl files) into ExportFolder.",
        "2. Data > Get Data > From Other Sources > Blank Query. Open Advanced Editor, paste the code from the",
        "   Queries sheet, click Done, and rename the query to the name in column A.",
        "   Do this for the queries in the order listed (top to bottom).",
        "3. For Elements, QuantitiesByType and Completeness choose Close & Load To > Table (or PivotTable Report).",
        "   RawElements, ExportFolder and PreferredFormat can stay as Connection only.",
        "4. Insert > PivotTable on QuantitiesByType: Rows = Category, Values = Count / Volume / Area.",
        "5. After a new conversion: Data > Refresh All.",
        "",
        "Values are summed as exported, no unit conversion. SQLite output is not read directly by Power Query;",
        "use the xlsx, csv or jsonl files that the converter writes next to it.",
        "",
        "Revit, AutoCAD and Excel are names of products of their respective owners and are used only to say what the",
        "tool works with. This kit is independent software and is not affiliated with or endorsed by them.",
    ]
    for i, t in enumerate(steps, 1):
        ws.cell(i, 1, t)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120

    cf = wb.create_sheet("Config")
    cf.append(["Setting", "Value"])
    cf.append(["ExportFolder", r"C:\DDC\exports\MyProject"])
    cf.append(["PreferredFormat", "xlsx"])
    t = Table(displayName="Config", ref="A1:B3")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    cf.add_table(t)
    cf.column_dimensions["A"].width = 20
    cf.column_dimensions["B"].width = 70
    cf["B2"].fill = PatternFill("solid", fgColor="FFEB9C")

    q = wb.create_sheet("Queries")
    q.append(["Query name", "What it does", "M code (paste into Advanced Editor)"])
    for c in q[1]:
        c.font = Font(bold=True)
    for name, src, why in ORDER:
        q.append([name, why, Path(src).read_text(encoding="utf-8")])
    q.column_dimensions["A"].width = 20
    q.column_dimensions["B"].width = 45
    q.column_dimensions["C"].width = 120
    for row in q.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=False, vertical="top")
        row[0].alignment = Alignment(vertical="top")
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(target)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "DDC_Excel_Kit.xlsx")
