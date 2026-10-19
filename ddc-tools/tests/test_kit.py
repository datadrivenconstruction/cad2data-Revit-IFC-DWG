import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "excel"))
import build_kit  # noqa: E402


def balanced(code):
    """Brackets balance outside strings and comments (cheap sanity check, not an M parser)."""
    pairs = {")": "(", "]": "[", "}": "{"}
    stack, i, n = [], 0, len(code)
    while i < n:
        ch = code[i]
        if code.startswith("//", i):
            i = code.find("\n", i)
            if i < 0:
                break
        elif ch == '"':
            i += 1
            while i < n and not (code[i] == '"' and code[i + 1:i + 2] != '"'):
                i += 2 if code[i] == '"' else 1
        elif ch in "([{":
            stack.append(ch)
        elif ch in pairs:
            if not stack or stack.pop() != pairs[ch]:
                return False
        i += 1
    return not stack


def test_m_files_are_balanced_and_let_in_paired():
    for p in list((ROOT / "powerbi").glob("*.pq")) + list((ROOT / "excel").glob("*.pq")):
        code = p.read_text(encoding="utf-8")
        assert balanced(code), p.name
        body = "\n".join(l for l in code.splitlines() if not l.strip().startswith("//"))
        if body.lstrip().startswith("let"):
            assert "\nin\n" in body or "\n    in\n" in body or "in\n" in body, p.name


def test_kit_workbook(tmp_path):
    out = tmp_path / "kit.xlsx"
    build_kit.build(str(out))
    wb = load_workbook(out)
    assert wb.sheetnames == ["Start here", "Config", "Queries"]
    assert "Config" in wb["Config"].tables
    names = [r[0].value for r in wb["Queries"].iter_rows(min_row=2)]
    assert names == ["ExportFolder", "PreferredFormat", "RawElements", "Elements", "QuantitiesByType", "Completeness"]
    assert "Folder.Files" in wb["Queries"]["C4"].value
