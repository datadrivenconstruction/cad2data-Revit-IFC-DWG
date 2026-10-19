"""Real converter output layouts: no manual column mapping needed."""
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ddc_tools import check, qto
from ddc_tools.cli import main
from ddc_tools.loader import load_ex

FIX = Path(__file__).parent / "fixtures"
IFC = FIX / "ifc_duplex_excerpt"
RVT = FIX / "rvt_properties"
ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def ifc_rows():
    return load_ex(str(IFC), fmt="jsonl")


@pytest.mark.parametrize("fmt", ["sqlite", "jsonl", "xlsx"])
def test_ifc_formats_give_identical_flat_tables(fmt, ifc_rows):
    other = load_ex(str(IFC), fmt=fmt)
    assert other.rows == ifc_rows.rows
    assert other.quantity_columns == ifc_rows.quantity_columns


def test_ifc_columns_are_understood(ifc_rows):
    rows = ifc_rows.rows
    assert len(rows) == 34                      # 41 objects minus 7 type objects
    from ddc_tools import columns as C
    cols = C.all_columns(rows)
    assert C.resolve(cols, "id") == "ElementId"
    assert C.resolve(cols, "guid") == "GlobalId"
    assert C.resolve(cols, "category") == "Category"
    assert C.resolve(cols, "type") == "TypeName"
    assert C.resolve(cols, "level") == "Level"
    wall = next(r for r in rows if r["Category"] == "IfcWallStandardCase")
    assert wall["TypeName"].startswith("Basic Wall:") and wall["Level"].startswith("Level")
    assert wall["Pset_WallCommon.IsExternal"] in (True, False)
    # an opening takes the storey of the wall it cuts
    assert all(r["Level"] for r in rows if r["Category"] == "IfcOpeningElement")
    # spatial structure objects have no storey; everything else does
    assert {r["Category"] for r in rows if not r["Level"]} == {"IfcProject", "IfcSite", "IfcBuilding"}


def test_ifc_qto_finds_quantities_and_levels(ifc_rows):
    res = qto.build_qto(ifc_rows.rows, quantity_hint=ifc_rows.quantity_columns, notes=ifc_rows.notes)
    q = res["quantity_columns"]
    assert "PSet_Revit_Dimensions.Volume" in q and "PSet_Revit_Dimensions.Area" in q
    assert "GSA Space Areas.GSA BIM Area" in q                    # stored IfcQuantityArea
    assert not any("Type_" in c for c in q)                        # type property sets are not summed
    assert res["columns_used"] == {"id": "ElementId", "category": "Category", "type": "TypeName", "level": "Level"}
    walls = next(r for r in res["by_category"] if r["category"] == "IfcWallStandardCase")
    assert walls["count"] == 6 and walls["PSet_Revit_Dimensions.Volume"] > 0
    spaces = next(r for r in res["by_category"] if r["category"] == "IfcSpace")
    assert spaces["GSA Space Areas.GSA BIM Area"] > 0
    assert {"Level 1", "Level 2"} <= {r["level"] for r in res["by_level"]}
    assert res["notes"]


def test_ifc_check_short_names_match_every_property_set(ifc_rows):
    spec = check.validate_spec({"rules": [
        {"id": "FIRE", "check": "pattern", "param": "FireRating", "regex": r"(EI|REI|E|R) ?\d{2,3}",
         "where": {"category": "IfcDoor"}},
        {"id": "GUID", "check": "unique", "param": "GlobalId"},
        {"id": "WALL", "check": "required", "param": "Pset_WallCommon.Reference",
         "where": {"category": ["IfcWall*"]}},
    ]})
    res = check.run_checks(ifc_rows.rows, spec)
    by = {r["id"]: r for r in res["rules"]}
    assert by["GUID"]["status"] == "PASS" and by["WALL"]["status"] == "PASS"
    assert by["FIRE"]["status"] == "FAIL" and by["FIRE"]["applicable"] == 4   # placeholder 'Fire Rating' text
    f = res["findings"][0]
    assert f["global_id"] and f["element_id"] and f["level"].startswith("Level")


def test_ifc_cli_with_example_rules(tmp_path):
    out = tmp_path / "c.xlsx"
    rc = main(["check", str(IFC), "--format", "sqlite", "--rules", str(ROOT / "examples" / "rules_ifc.yaml"),
               "--xlsx", str(out), "-q"])
    assert rc == 1                       # door fire rating placeholders
    assert load_workbook(out).sheetnames == ["Summary", "Details"]
    q = tmp_path / "q.xlsx"
    assert main(["qto", str(IFC), "--format", "xlsx", "--xlsx", str(q), "-q"]) == 0
    ws = load_workbook(q)["Read me"]
    assert any(c.value == "Source note" for c in ws["A"])


def test_rvt_saved_properties():
    ld = load_ex(str(RVT))
    assert ld.files == ["properties.jsonl"] and len(ld.rows) == 4
    by = {r["ElementId"]: r for r in ld.rows}
    assert by["100"]["Category"] == "-2000011" and by["100"]["Mark"] == "W-1" and by["100"]["Length"] == 12.5
    assert by["101"]["Base Level"] == 5000          # element reference resolved to its id
    assert "Comments" not in by["102"]              # unresolved named parameters are not invented
    assert any("category is a numeric id" in n for n in ld.notes)
    res = qto.build_qto(ld.rows, notes=ld.notes)
    walls = next(r for r in res["by_category"] if r["category"] == "-2000011")
    assert walls["count"] == 2 and walls["Length"] == 19.75
    spec = check.validate_spec({"rules": [{"id": "U", "check": "unique", "param": "Mark"}]})
    assert check.run_checks(ld.rows, spec)["summary"]["fail"] == 1   # W-1 twice
