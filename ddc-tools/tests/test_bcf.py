"""Tests for ``ddc bcf`` and the ``--bcf`` options. Files are validated against the official buildingSMART
BCF XSD (downloaded on demand, see conftest.py) for BCF 2.1 and 3.0."""
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ddc_tools import bcf, cli, diff

FIX = Path(__file__).parent / "fixtures"
ROOT = Path(__file__).resolve().parents[1]
A, B = FIX / "ifc_duplex_excerpt", FIX / "ifc_duplex_excerpt_v2"
IFC_RULES = ROOT / "examples" / "rules_ifc.yaml"


@pytest.fixture(scope="module")
def check_json(tmp_path_factory):
    p = tmp_path_factory.mktemp("chk") / "check.json"
    rc = cli.main(["check", str(A), "--rules", str(IFC_RULES), "--json", str(p), "-q"])
    assert rc == 1
    return p


@pytest.fixture(scope="module")
def diff_json(tmp_path_factory):
    p = tmp_path_factory.mktemp("dif") / "diff.json"
    assert cli.main(["diff", str(A), str(B), "--json", str(p), "-q"]) == 1
    return p


# ---------------------------------------------------------------------------- schema validity

@pytest.mark.parametrize("version", ["2.1", "3.0"])
@pytest.mark.parametrize("per_element", [False, True])
def test_check_export_is_valid_bcf(check_json, tmp_path, bcf_validator, version, per_element):
    out = tmp_path / "c.bcf"
    args = ["bcf", "export", str(check_json), "-o", str(out), "--version", version]
    assert cli.main(args + (["--per-element"] if per_element else [])) == 0
    errors, kinds = bcf_validator(out, version)
    assert errors == []
    assert {"markup", "visinfo", "version", "project"} <= kinds and (("extensions" in kinds) == (version == "3.0"))


@pytest.mark.parametrize("version", ["2.1", "3.0"])
@pytest.mark.parametrize("per_element", [False, True])
def test_diff_export_is_valid_bcf(diff_json, tmp_path, bcf_validator, version, per_element):
    out = tmp_path / "d.bcf"
    args = ["bcf", "export", str(diff_json), "-o", str(out), "--version", version]
    assert cli.main(args + (["--per-element"] if per_element else [])) == 0
    errors, _ = bcf_validator(out, version)
    assert errors == []


@pytest.mark.parametrize("version", ["2.1", "3.0"])
def test_options_stay_valid(diff_json, tmp_path, bcf_validator, version):
    out = tmp_path / "o.bcf"
    rc = cli.main(["bcf", "export", str(diff_json), "-o", str(out), "--only", "added", "--camera", "10,-10,8,-1,1,-0.5",
                   "--author", "me@example.org", "--model-name", "Duplex_v2.ifc", "--max-components", "1",
                   "--version", version])
    assert rc == 0
    assert bcf_validator(out, version)[0] == []
    res = bcf.read_bcf(str(out))
    assert sorted(t["title"] for t in res["topics"]) == ["Added: IfcDoor (1)", "Added: IfcWallStandardCase (1)"]
    assert all(t["author"] == "me@example.org" and t["files"] == ["Duplex_v2.ifc"] for t in res["topics"])
    assert all(v["camera"] == "perspective" for t in res["topics"] for v in t["viewpoints"])


# ---------------------------------------------------------------------------- content

def test_check_topics_follow_rules_and_select_ifc_guids(check_json, tmp_path):
    rep = json.loads(check_json.read_text(encoding="utf-8"))
    out = tmp_path / "c.bcf"
    res = bcf.export_report(rep, str(out))
    failing = {f["rule_id"] for f in rep["findings"]}
    assert res["kind"] == "check" and res["topics"] == len(failing) == 2
    back = bcf.read_bcf(str(out))
    assert back["version"] == "2.1" and len(back["topics"]) == 2
    guids_in_bcf = {c["IfcGuid"] for t in back["topics"] for v in t["viewpoints"] for c in v["selection"]}
    guids_in_report = {f["global_id"] for f in rep["findings"] if f["global_id"]}
    assert guids_in_bcf == guids_in_report and len(guids_in_report) >= 4
    door = next(t for t in back["topics"] if "DOOR-FIRE" in t["title"])
    assert door["priority"] == "High" and door["status"] == "Open" and door["type"] == "Error"
    assert "ddc-check" in door["labels"] and "IFC-DOOR-FIRE-VALUE" in door["labels"]
    assert "Fire Rating" in door["description"] or "FireRating" in door["description"]
    assert len(door["files"]) == 1


def test_diff_topics_group_by_status_and_category_and_colour(diff_json, tmp_path):
    rep = json.loads(diff_json.read_text(encoding="utf-8"))
    out = tmp_path / "d.bcf"
    res = bcf.export_report(rep, str(out))
    back = bcf.read_bcf(str(out))
    titles = sorted(t["title"] for t in back["topics"])
    assert titles == ["Added: IfcDoor (1)", "Added: IfcWallStandardCase (1)", "Modified: IfcSlab (1)",
                      "Modified: IfcSpace (1)", "Modified: IfcWallStandardCase (1)", "Modified: IfcWindow (1)",
                      "Removed: IfcDoor (1)", "Removed: IfcFurnishingElement (1)"]
    assert res["components"] == 8
    wall = next(t for t in back["topics"] if t["title"] == "Modified: IfcWallStandardCase (1)")
    assert "Length: 3.7915 -> 4.2915 (+0.5)" in wall["description"] and "FireRating: (empty) -> EI 60" in wall["description"]
    vp = wall["viewpoints"][0]
    assert [c["IfcGuid"] for c in vp["selection"]] == ["2O2Fr$t4X7Zf8NOew3FNld"]
    assert vp["coloring"][0]["color"] == "FA9E1A" and vp["default_visibility"] == "true"
    added = next(t for t in back["topics"] if t["title"] == "Added: IfcDoor (1)")
    assert added["viewpoints"][0]["coloring"][0]["color"] == "2EB84D"
    removed = next(t for t in back["topics"] if t["title"] == "Removed: IfcDoor (1)")
    assert removed["priority"] == "High" and removed["viewpoints"][0]["coloring"][0]["color"] == "DB3232"


def test_per_element_topics(diff_json, tmp_path):
    rep = json.loads(diff_json.read_text(encoding="utf-8"))
    res = bcf.export_report(rep, str(tmp_path / "p.bcf"), per_element=True)
    assert res["topics"] == 8 and res["viewpoints"] == 8


def test_ids_without_ifc_guid_use_authoring_tool_id(tmp_path):
    rep = {"tool": "ddc diff", "a": {"source": "A"}, "b": {"source": "B.converted"}, "elements": [
        {"status": "modified", "key": "2567", "global_id": "", "element_id": "2567", "category": "OST_Walls",
         "type": "Wall 1", "level": "Level 1", "name": "", "changes": [
             {"field": "Mark", "kind": "property", "old": "W-1", "new": "W-2"}]},
        {"status": "added", "key": "9", "global_id": "not-a-guid", "element_id": "9", "category": "OST_Walls",
         "type": "", "level": "", "name": "", "changes": []}]}
    out = tmp_path / "r.bcf"
    bcf.export_report(rep, str(out), version="3.0")
    back = bcf.read_bcf(str(out))
    comps = [c for t in back["topics"] for v in t["viewpoints"] for c in v["selection"]]
    assert {(c["IfcGuid"], c["AuthoringToolId"]) for c in comps} == {("", "2567"), ("", "9")}
    assert all(c["OriginatingSystem"] for c in comps)
    assert back["version"] == "3.0" and back["topics"][0]["files"] == ["B.converted"]


def test_max_components_truncates_and_says_so(tmp_path):
    findings = [{"rule_id": "R1", "severity": "warning", "element_id": str(i), "global_id": "", "category": "C",
                 "type": "", "level": "", "parameter": "p", "value": "", "message": f"p is empty ({i})"} for i in range(30)]
    rep = {"tool": "ddc check", "ruleset": "x", "source": "s", "rules": [{"id": "R1", "title": "p set", "check": "required",
                                                                           "severity": "warning"}], "findings": findings}
    out = tmp_path / "t.bcf"
    res = bcf.export_report(rep, str(out), max_components=10)
    assert res["components"] == 10
    t = bcf.read_bcf(str(out))["topics"][0]
    assert "20 more element(s)" in t["description"] and t["priority"] == "Normal"


def test_wrong_report_and_bad_options(tmp_path):
    p = tmp_path / "x.json"
    p.write_text('{"tool": "something else"}', encoding="utf-8")
    assert cli.main(["bcf", "export", str(p), "-o", str(tmp_path / "o.bcf")]) == 2
    p.write_text("not json", encoding="utf-8")
    assert cli.main(["bcf", "export", str(p), "-o", str(tmp_path / "o.bcf")]) == 2
    with pytest.raises(bcf.BcfError):
        bcf.parse_camera("1,2,3")


def test_empty_report_gives_valid_empty_file(tmp_path, bcf_validator):
    rep = {"tool": "ddc check", "ruleset": "x", "source": "s", "rules": [], "findings": []}
    for v in ("2.1", "3.0"):
        out = tmp_path / f"e{v}.bcf"
        assert bcf.export_report(rep, str(out), version=v)["topics"] == 0
        assert bcf_validator(out, v)[0] == []
        assert bcf.read_bcf(str(out))["topics"] == []


def test_check_and_diff_bcf_shortcuts(tmp_path, capsys):
    c, d = tmp_path / "c.bcf", tmp_path / "d.bcf"
    assert cli.main(["check", str(A), "--rules", str(IFC_RULES), "--bcf", str(c), "--bcf-version", "3.0"]) == 1
    assert "BCF 3.0: 2 topics" in capsys.readouterr().out
    assert cli.main(["diff", str(A), str(B), "--bcf", str(d), "--bcf-per-element"]) == 1
    assert "BCF 2.1: 8 topics" in capsys.readouterr().out
    assert bcf.read_bcf(str(c))["version"] == "3.0" and len(bcf.read_bcf(str(d))["topics"]) == 8


def test_export_is_reproducible_apart_from_the_date(diff_json, tmp_path):
    rep = json.loads(diff_json.read_text(encoding="utf-8"))
    r1, r2 = tmp_path / "1.bcf", tmp_path / "2.bcf"
    bcf.export_report(rep, str(r1))
    bcf.export_report(rep, str(r2))
    t1 = {t["guid"]: t["title"] for t in bcf.read_bcf(str(r1))["topics"]}
    t2 = {t["guid"]: t["title"] for t in bcf.read_bcf(str(r2))["topics"]}
    assert t1 == t2


# ---------------------------------------------------------------------------- reader

def _zip(path, files):
    with zipfile.ZipFile(path, "w") as z:
        for n, d in files.items():
            z.writestr(n, d)


HAND_MARKUP = """<?xml version="1.0" encoding="utf-8"?>
<Markup xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:xsd="http://www.w3.org/2001/XMLSchema">
  <Header><File IfcProject="0J$yPqHBD0JrHN3Z4Sc9f2" isExternal="false"><Filename>House.ifc</Filename></File></Header>
  <Topic Guid="aaaaaaaa-0000-4000-8000-000000000001" TopicType="Clash" TopicStatus="Active">
    <Title>Wall clashes with duct</Title><Priority>Critical</Priority><Labels>MEP</Labels><Labels>Level 2</Labels>
    <CreationDate>2025-03-01T10:00:00+01:00</CreationDate><CreationAuthor>a@example.org</CreationAuthor>
    <AssignedTo>b@example.org</AssignedTo><Description>Move the duct</Description>
  </Topic>
  <Comment Guid="bbbbbbbb-0000-4000-8000-000000000001"><Date>2025-03-02T09:00:00Z</Date><Author>b@example.org</Author>
    <Comment>Will do</Comment><Viewpoint Guid="cccccccc-0000-4000-8000-000000000001"/></Comment>
  <Viewpoints Guid="cccccccc-0000-4000-8000-000000000001"><Viewpoint>vp.bcfv</Viewpoint><Snapshot>snap.png</Snapshot></Viewpoints>
  <Viewpoints Guid="dddddddd-0000-4000-8000-000000000001"><Viewpoint>missing.bcfv</Viewpoint></Viewpoints>
</Markup>"""
HAND_VP = """<?xml version="1.0"?>
<VisualizationInfo Guid="cccccccc-0000-4000-8000-000000000001">
  <Components><Selection><Component IfcGuid="0J$yPqHBD0JrHN3Z4Sc9f2"><AuthoringToolId>123</AuthoringToolId></Component></Selection>
  <Visibility DefaultVisibility="false"><Exceptions><Component IfcGuid="1hOSvn6df7F8_7GcBWlR72"/></Exceptions></Visibility></Components>
  <OrthogonalCamera><CameraViewPoint><X>0</X><Y>0</Y><Z>0</Z></CameraViewPoint></OrthogonalCamera>
</VisualizationInfo>"""


def test_read_hand_written_bcf21(tmp_path):
    p = tmp_path / "h.bcf"
    _zip(p, {"bcf.version": '<Version VersionId="2.1"><DetailedVersion>2.1</DetailedVersion></Version>',
             "p/markup.bcf": HAND_MARKUP, "p/vp.bcfv": HAND_VP, "p/snap.png": b"x"})
    r = bcf.read_bcf(str(p))
    t = r["topics"][0]
    assert (t["title"], t["type"], t["status"], t["priority"], t["assigned_to"]) == \
        ("Wall clashes with duct", "Clash", "Active", "Critical", "b@example.org")
    assert t["labels"] == ["MEP", "Level 2"] and t["files"] == ["House.ifc"]
    assert t["comments"][0]["text"] == "Will do" and t["comments"][0]["viewpoint"].startswith("cccccccc")
    v0, v1 = t["viewpoints"]
    assert v0["camera"] == "orthogonal" and v0["default_visibility"] == "false" and v0["snapshot"] == "snap.png"
    assert v0["selection"][0]["AuthoringToolId"] == "123" and v0["visible_exceptions"][0]["IfcGuid"] == "1hOSvn6df7F8_7GcBWlR72"
    assert v1["error"] == "viewpoint file missing"
    assert r["summary"]["topics"] == 1 and r["summary"]["comments"] == 1 and r["summary"]["by_status"] == {"Active": 1}


def test_read_to_xlsx_and_json(diff_json, tmp_path):
    rep = json.loads(diff_json.read_text(encoding="utf-8"))
    b = tmp_path / "d.bcf"
    bcf.export_report(rep, str(b), version="3.0")
    x, j = tmp_path / "r.xlsx", tmp_path / "r.json"
    assert cli.main(["bcf", "read", str(b), "--xlsx", str(x), "--json", str(j), "-q"]) == 0
    wb = load_workbook(x)
    assert wb.sheetnames == ["Summary", "Topics", "Comments", "Components"]
    topics = [[c.value for c in r] for r in wb["Topics"].iter_rows(min_row=1)]
    assert topics[0][:4] == ["Title", "Status", "Priority", "Type"] and len(topics) == 9
    comps = [[c.value for c in r] for r in wb["Components"].iter_rows(min_row=2)]
    assert {r[2] for r in comps} == {"selected", "colored"} and len(comps) == 16
    assert "2O2Fr$t4X7Zf8NOew3FNld" in {r[3] for r in comps}
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=1) if r[0].value}
    assert summary["BCF version"] == "3.0" and summary["Topics"] == 8 and summary["Open"] == 8
    assert json.loads(j.read_text(encoding="utf-8"))["summary"]["topics"] == 8


def test_reader_rejects_hostile_or_broken_files(tmp_path):
    bomb = tmp_path / "bomb.bcf"
    _zip(bomb, {"bcf.version": '<!DOCTYPE v [<!ENTITY a "aaaa">]><Version VersionId="2.1">&a;</Version>'})
    with pytest.raises(bcf.BcfError, match="DOCTYPE"):
        bcf.read_bcf(str(bomb))
    junk = tmp_path / "junk.bcf"
    junk.write_bytes(b"this is not a zip")
    with pytest.raises(bcf.BcfError):
        bcf.read_bcf(str(junk))
    broken = tmp_path / "broken.bcf"
    _zip(broken, {"bcf.version": "<Version", "x/markup.bcf": "<Markup>"})
    with pytest.raises(bcf.BcfError, match="invalid XML"):
        bcf.read_bcf(str(broken))
    assert cli.main(["bcf", "read", str(junk)]) == 2
    assert cli.main(["bcf", "read", str(tmp_path / "missing.bcf")]) == 2


def test_reader_ignores_nested_non_topic_folders(tmp_path):
    p = tmp_path / "n.bcf"
    _zip(p, {"bcf.version": '<Version VersionId="2.1"/>', "a/b/markup.bcf": HAND_MARKUP})
    assert bcf.read_bcf(str(p))["topics"] == []
