import json

import pytest
from openpyxl import load_workbook

from ddc_tools import check
from ddc_tools.cli import main


def spec(*rules):
    return check.validate_spec({"rules": list(rules)})


def run(rows, *rules):
    return check.run_checks(rows, spec(*rules))


def rule(res, rid):
    return next(r for r in res["rules"] if r["id"] == rid)


def test_required_by_category(rows):
    res = run(rows, {"id": "A", "check": "required", "param": "Fire Rating", "where": {"category": "Walls"}})
    r = rule(res, "A")
    assert (r["applicable"], r["failed_elements"], r["status"]) == (9, 2, "FAIL")
    f = res["findings"][0]
    assert f["element_id"] and f["global_id"].startswith("3Xk") and f["category"] == "Walls"


def test_required_missing_column_fails_everything(rows):
    r = rule(run(rows, {"id": "A", "check": "required", "param": "Nope"}), "A")
    assert r["failed_elements"] == len(rows) and "not found" in r["note"]


def test_not_empty_missing_column_is_skipped(rows):
    r = rule(run(rows, {"id": "A", "check": "not_empty", "param": "Nope"}), "A")
    assert r["status"] == "SKIP"


def test_pattern_uses_full_match_and_ignores_empty(rows):
    res = run(rows, {"id": "A", "check": "pattern", "param": "Mark", "regex": r"W-\d{3}",
                     "where": {"category": "Walls"}})
    assert rule(res, "A")["failed_elements"] == 1
    assert res["findings"][0]["value"] == "wall7"


def test_unique(rows):
    res = run(rows, {"id": "A", "check": "unique", "param": "Mark"})
    assert rule(res, "A")["failed_elements"] == 2
    assert {f["value"] for f in res["findings"]} == {"W-001"}


def test_unique_scope(rows):
    res = run(rows, {"id": "A", "check": "unique", "param": "Mark", "scope": "Level"})
    # the duplicate W-001 sits on Level 2 while the other is on Level 1
    assert rule(res, "A")["status"] == "PASS"


def test_range_and_non_numeric(rows):
    res = run(rows, {"id": "A", "check": "range", "param": "Volume", "min": 0.1})
    msgs = [f["message"] for f in res["findings"]]
    assert any("not a number" in m for m in msgs) and any("outside" in m for m in msgs)


def test_enum_case_insensitive(rows):
    res = run(rows, {"id": "A", "check": "enum", "param": "Fire Rating", "values": ["ei 30", "EI 60", "R 90"],
                     "severity": "warning"})
    assert rule(res, "A")["status"] == "PASS"
    res = run(rows, {"id": "A", "check": "enum", "param": "Fire Rating", "values": ["EI 60"],
                     "severity": "warning"})
    assert rule(res, "A")["status"] == "WARN"


def test_severity_decides_status_and_exit(rows):
    res = run(rows, {"id": "A", "check": "not_empty", "param": "Level", "severity": "warning"})
    assert res["summary"]["overall"] == "WARN"
    assert check.exit_code(res) == 0 and check.exit_code(res, "warning") == 1


def test_no_match_is_skip(rows):
    r = rule(run(rows, {"id": "A", "check": "required", "param": "Mark", "where": {"category": "Roofs"}}), "A")
    assert r["status"] == "SKIP" and r["applicable"] == 0


def test_where_regex_and_glob(rows):
    r = rule(run(rows, {"id": "A", "check": "required", "param": "Mark", "where": {"category_regex": "^Gen"}}), "A")
    assert r["applicable"] == 1 and r["failed_elements"] == 1
    r = rule(run(rows, {"id": "A", "check": "required", "param": "Mark", "where": {"type": "*200mm"}}), "A")
    assert r["applicable"] > 0


@pytest.mark.parametrize("bad", [
    {"rules": []},
    {"rules": [{"check": "nope", "param": "x"}]},
    {"rules": [{"check": "pattern", "param": "x"}]},
    {"rules": [{"check": "pattern", "param": "x", "regex": "("}]},
    {"rules": [{"check": "range", "param": "x"}]},
    {"rules": [{"check": "required"}]},
    {"rules": [{"check": "required", "param": "x", "severity": "fatal"}]},
    {"rules": [{"id": "a", "check": "required", "param": "x"}, {"id": "a", "check": "required", "param": "x"}]},
    {"rules": [{"check": "required", "param": "x", "where": {"colour": "red"}}]},
])
def test_invalid_rules_rejected(bad):
    with pytest.raises(check.RuleError):
        check.validate_spec(bad)


def test_ifc_style_columns():
    rows = [{"GlobalId": "g1", "IfcClass": "IfcWall", "Pset_WallCommon.FireRating": None},
            {"GlobalId": "g2", "IfcClass": "IfcWall", "Pset_WallCommon.FireRating": "EI 60"}]
    res = run(rows, {"id": "A", "check": "required", "param": "FireRating", "where": {"category": "IfcWall"}})
    assert res["findings"][0]["global_id"] == "g1"


def test_cli_end_to_end(export_dir, rules_basic, tmp_path):
    x, j = tmp_path / "r" / "check.xlsx", tmp_path / "check.json"
    rc = main(["check", str(export_dir), "--rules", str(rules_basic), "--xlsx", str(x), "--json", str(j), "-q"])
    assert rc == 1
    data = json.loads(j.read_text(encoding="utf-8"))
    assert data["summary"]["overall"] == "FAIL" and data["findings"]
    wb = load_workbook(x)
    assert wb.sheetnames == ["Summary", "Details"]
    statuses = {c.value for c in wb["Summary"]["A"] if c.value in ("PASS", "FAIL", "WARN", "SKIP", "INFO")}
    assert {"FAIL", "PASS"} <= statuses
    hdr = [c.value for c in wb["Details"][1]]
    assert "ElementId" in hdr and "GlobalId" in hdr


def test_cli_exit_codes(export_dir, tmp_path, capsys):
    ok = tmp_path / "ok.yaml"
    ok.write_text("rules:\n  - {id: X, check: not_empty, param: Category}\n", encoding="utf-8")
    assert main(["check", str(export_dir), "--rules", str(ok), "-q"]) == 0
    assert main(["check", str(export_dir), "--rules", str(tmp_path / "missing.yaml"), "-q"]) == 2
    assert main(["check", str(tmp_path / "nowhere"), "--rules", str(ok), "-q"]) == 2
    assert "error" in capsys.readouterr().err
