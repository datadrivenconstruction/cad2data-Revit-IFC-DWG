"""Tests for ``ddc diff`` on a real IFC to Excel excerpt and an edited copy of it (tests/fixtures/README.md)."""
import copy
import json
import shutil
import struct
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ddc_tools import cli, diff, glbdiff, table
from _glbutil import write_cubes

FIX = Path(__file__).parent / "fixtures"
DATA = Path(__file__).parent / "data"
A = FIX / "ifc_duplex_excerpt"
B = FIX / "ifc_duplex_excerpt_v2"

WALL = "2O2Fr$t4X7Zf8NOew3FNld"          # modified: re-measured, FireRating added, now external
WINDOW = "1hOSvn6df7F8_7GcBWlRBU"        # modified: renamed
SLAB = "2O2Fr$t4X7Zf8NOew3FK4F"          # modified: thicker
SPACE = "0BTBFw6f90Nfh9rP1dlXrr"         # modified: stored area quantity
NEW_WALL, NEW_DOOR = "3Zk0a1dXj9ZwB2Q7Ls4tDx", "1hOSvn6df7F8_7GcBWlNEW"
GONE_DOOR, GONE_CABINET = "1hOSvn6df7F8_7GcBWlS9F", "0wkEuT1wr1kOyafLY4vy3H"
ROOF = "0jf0rYHfX3RAB3bSIRjmxl"          # STEP id renumbered in version 2, GlobalId kept


@pytest.fixture(scope="module")
def res():
    return diff.diff_paths(str(A), str(B))


def el(res, guid):
    return next(e for e in res["elements"] if e["global_id"] == guid)


def ch(e, field):
    return next(c for c in e["changes"] if c["field"] == field)


# ---------------------------------------------------------------------------- matching and statuses

def test_summary_and_key(res):
    s = res["summary"]
    assert (s["added"], s["removed"], s["modified"], s["unchanged"]) == (2, 2, 4, 18)
    assert s["total_a"] == s["total_b"] == 24 and not s["identical"]
    assert res["key"]["column"] == "GlobalId" and res["key"]["method"] == "key:GlobalId"
    # two elements got a new STEP id: reported, not counted as a change
    assert len(res["warnings"]) == 1 and "2 matched elements have a different ID" in res["warnings"][0]


def test_added_and_removed(res):
    assert {e["global_id"] for e in res["elements"] if e["status"] == "added"} == {NEW_WALL, NEW_DOOR}
    assert {e["global_id"] for e in res["elements"] if e["status"] == "removed"} == {GONE_DOOR, GONE_CABINET}
    assert el(res, NEW_WALL)["category"] == "IfcWallStandardCase"
    assert el(res, GONE_CABINET)["a"]["ID"] == "13331" and el(res, GONE_CABINET)["b"] == {}


def test_quantity_changes_have_old_new_delta(res):
    e = el(res, WALL)
    assert e["status"] == "modified"
    c = ch(e, "Length")
    assert c["kind"] == "quantity"
    assert (c["old"], c["new"]) == (pytest.approx(3.7915), pytest.approx(4.2915))
    assert c["delta"] == pytest.approx(0.5) and c["delta_pct"] == pytest.approx(13.19, abs=0.01)
    assert ch(e, "Area")["new"] == pytest.approx(11.9955)
    slab = el(res, SLAB)
    assert ch(slab, "Width")["delta"] == pytest.approx(0.023) and ch(slab, "Volume")["kind"] == "quantity"
    assert ch(el(res, SPACE), "Area")["new"] == 19.25 and ch(el(res, SPACE), "Volume")["new"] == 50.05


def test_property_changes_old_to_new(res):
    e = el(res, WALL)
    fr = ch(e, "FireRating")
    assert fr["kind"] == "property" and fr["old"] == "" and fr["new"] == "EI 60"
    ext = ch(e, "IsExternal")
    assert (ext["old"], ext["new"]) == (False, True)
    assert ch(e, "Top Offset")["new"] == pytest.approx(-0.25)
    assert ch(el(res, SLAB), "Room Bounding")["new"] is False


def test_attribute_change_is_not_a_property(res):
    c = ch(el(res, WINDOW), "Name")
    assert c["kind"] == "attribute" and c["new"].endswith("(rev B)") and c["old"] + " (rev B)" == c["new"]
    assert len(el(res, WINDOW)["changes"]) == 1


def test_step_id_churn_is_not_a_change(res):
    e = el(res, ROOF)
    assert e["status"] == "unchanged"
    assert e["a"]["ID"] == "22475" and e["b"]["ID"] == "122475"


def test_by_category_and_quantity_totals(res):
    cats = {c["category"]: c for c in res["by_category"]}
    assert cats["IfcDoor"]["added"] == 1 and cats["IfcDoor"]["removed"] == 1
    assert cats["IfcWallStandardCase"]["modified"] == 1 and cats["IfcWallStandardCase"]["added"] == 1
    assert cats["IfcWallStandardCase"]["total_a"] == 6 and cats["IfcWallStandardCase"]["total_b"] == 7
    q = next(x for x in res["quantities"] if x["category"] == "IfcSlab" and x["quantity"] == "Volume")
    assert q["delta"] == pytest.approx(q["b"] - q["a"]) and q["delta"] > 0
    assert any(f["field"] == "Length" for f in res["changed_fields"])


def test_identical_inputs_and_other_formats():
    # the same export read from jsonl, sqlite and xlsx must give no differences at all
    j, s, x = A / "data.jsonl", A / "data.sqlite", A / "drawing.xlsx"
    for left, right in ((j, s), (s, x), (x, str(A))):
        r = diff.diff_paths(str(left), str(right))
        assert r["summary"]["identical"], (left, right, r["summary"])
        assert r["summary"]["unchanged"] == 24


def test_unchanged_elements_in_json_only_on_request(res):
    assert all(e["status"] != "unchanged" for e in diff.public_result(res)["elements"])
    assert len(diff.public_result(res, True)["elements"]) == 26
    json.dumps(diff.public_result(res, True))


# ---------------------------------------------------------------------------- tolerance, ignore, key, fingerprint

def _tables():
    return table.build_table(str(A)), table.build_table(str(A))


def test_float_noise_is_not_a_change_but_a_real_edit_is():
    ta, tb = _tables()
    w = next(r for r in tb.rows if r["GlobalId"] == WALL)
    w["Length"] = w["Length"] * (1 + 1e-9)
    assert diff.diff_tables(ta, tb)["summary"]["identical"]
    w["Length"] = w["Length"] * 1.001
    r = diff.diff_tables(ta, tb)
    assert r["summary"]["modified"] == 1
    assert diff.diff_tables(ta, tb, tol=0.01)["summary"]["identical"]


def test_ignore_columns():
    ta, tb = _tables()
    w = next(r for r in tb.rows if r["GlobalId"] == WALL)
    w["Name"] = "renamed"
    assert diff.diff_tables(ta, tb)["summary"]["modified"] == 1
    assert diff.diff_tables(ta, tb, ignore=["Nam*"])["summary"]["identical"]


def test_fingerprint_matches_elements_with_regenerated_ids():
    ta, tb = _tables()
    for n, r in enumerate(tb.rows):                     # every id is new in B
        r["GlobalId"] = f"NEW{n:019d}"
        r["ID"] = 500000 + n
    r = diff.diff_tables(ta, tb)
    s = r["summary"]
    assert r["key"]["method"] == "fingerprint"
    assert s["unchanged"] == 24 and s["matched_by_fingerprint"] == 24 and s["identical"]
    assert any("matched by fingerprint" in w or "almost no" in w for w in r["warnings"])
    assert all(e["matched_by"] == "fingerprint" for e in r["elements"])
    with pytest.raises(diff.DiffError):
        diff.diff_tables(ta, tb, fingerprint=False)


def test_fingerprint_does_not_hide_a_real_edit():
    ta, tb = _tables()
    for n, r in enumerate(tb.rows):
        r["GlobalId"] = f"NEW{n:019d}"
        r["ID"] = 500000 + n
    next(r for r in tb.rows if r["Category"] == "IfcRoof")["Volume"] = 99.0
    r = diff.diff_tables(ta, tb)
    changed = [(e["category"], e["status"]) for e in r["elements"] if e["status"] != "unchanged"]
    assert sorted(changed) == [("IfcRoof", "added"), ("IfcRoof", "removed")]       # honest: not guessed as "modified"
    assert r["summary"]["matched_by_fingerprint"] == 23


def test_key_falls_back_when_the_first_candidate_does_not_overlap():
    ta, tb = table.build_table(str(DATA / "rvt_converted")), table.build_table(str(DATA / "rvt_converted"))
    for r in tb.rows:
        if r.get("GlobalId"):
            r["GlobalId"] = r["GlobalId"][::-1]          # IFC GlobalIds regenerated, ElementIds kept
    r = diff.diff_tables(ta, tb)
    assert r["key"]["column"] == "ElementId"
    assert r["summary"]["unchanged"] == 4 and r["summary"]["modified"] == 0
    assert any("different GlobalId" in w for w in r["warnings"])        # reported, not counted as a change
    assert diff.diff_tables(ta, tb, key="ElementId")["key"]["column"] == "ElementId"
    with pytest.raises(diff.DiffError):
        diff.diff_tables(ta, tb, key="NoSuchColumn")


def test_duplicate_keys_are_paired_in_order():
    ta, tb = _tables()
    ta.rows.append(copy.deepcopy(ta.rows[3]))
    tb.rows.append(copy.deepcopy(tb.rows[3]))
    r = diff.diff_tables(ta, tb)
    assert r["summary"]["identical"] and any("repeat" in w for w in r["warnings"])


def test_empty_inputs_raise():
    ta, tb = _tables()
    ta.rows, tb.rows = [], []
    with pytest.raises(diff.DiffError):
        diff.diff_tables(ta, tb)


def test_parse_fail_on():
    assert diff.parse_fail_on("any") == ["added", "removed", "modified"]
    assert diff.parse_fail_on("none") == [] and diff.parse_fail_on("removed, modified") == ["removed", "modified"]
    with pytest.raises(diff.DiffError):
        diff.parse_fail_on("deleted")


# ---------------------------------------------------------------------------- DWG

def test_dwg_diff_by_handle(tmp_path):
    src = DATA / "dwg" / "sample.jsonl"
    lines = src.read_text(encoding="utf-8").splitlines()
    out = []
    moved = removed = None
    for ln in lines:
        d = json.loads(ln)
        if d.get("type_name") == "LINE" and moved is None:
            d["x2"] = float(d["x2"]) + 5.0
            moved = d["handle"]
        elif d.get("type_name") == "CIRCLE" and removed is None:
            removed = d["handle"]
            continue
        out.append(json.dumps(d))
    p = tmp_path / "v2.jsonl"
    p.write_text("\n".join(out) + "\n", encoding="utf-8")
    r = diff.diff_paths(str(src), str(p))
    assert r["key"]["column"] == "ID"
    by = {e["key"]: e for e in r["elements"] if e["status"] != "unchanged"}
    assert by[moved]["status"] == "modified" and by[removed]["status"] == "removed"
    assert any(c["field"] == "Length" and c["kind"] == "quantity" for c in by[moved]["changes"])


# ---------------------------------------------------------------------------- XLSX

def test_xlsx_changelog(res, tmp_path):
    p = tmp_path / "changes.xlsx"
    diff.write_xlsx(res, str(p))
    wb = load_workbook(p)
    assert wb.sheetnames == ["Summary", "Changelog", "Elements", "By category", "Quantities"]
    ws = wb["Changelog"]
    head = [c.value for c in ws[1]]
    assert head[:3] == ["Status", "Key", "GlobalId"] and "Old" in head and "New" in head and "Delta" in head
    rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    length = next(r for r in rows if r["GlobalId"] == WALL and r["Field"] == "Length")
    assert length["Status"] == "modified" and length["Change kind"] == "quantity"
    assert length["Old"] == pytest.approx(3.7915) and length["New"] == pytest.approx(4.2915)
    assert length["Delta"] == pytest.approx(0.5)
    fire = next(r for r in rows if r["GlobalId"] == WALL and r["Field"] == "FireRating")
    assert fire["Old"] in (None, "") and fire["New"] == "EI 60"
    assert {r["Status"] for r in rows} == {"added", "removed", "modified"}
    assert ws.cell(2, 1).fill.fgColor.rgb.endswith("C6EFCE")             # added = green
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=1) if r[0].value}
    assert summary["added"] == 2 and summary["removed"] == 2 and summary["modified"] == 4 and summary["unchanged"] == 18
    assert all(r[0].value != "unchanged" for r in wb["Elements"].iter_rows(min_row=2))
    cats = {r[0].value: [c.value for c in r] for r in wb["By category"].iter_rows(min_row=2)}
    assert cats["IfcDoor"][1:3] == [1, 1]


def test_xlsx_with_unchanged(res, tmp_path):
    p = tmp_path / "all.xlsx"
    diff.write_xlsx(res, str(p), include_unchanged=True)
    assert load_workbook(p)["Elements"].max_row == 27


# ---------------------------------------------------------------------------- CLI and exit codes

def test_cli_exit_codes(tmp_path, capsys):
    x, j = tmp_path / "d.xlsx", tmp_path / "d.json"
    assert cli.main(["diff", str(A), str(B), "--xlsx", str(x), "--json", str(j)]) == 1
    out = capsys.readouterr().out
    assert "2 added, 2 removed, 4 modified, 18 unchanged" in out and "key:GlobalId" in out
    data = json.loads(j.read_text(encoding="utf-8"))
    assert data["tool"] == "ddc diff" and data["summary"]["modified"] == 4 and x.exists()
    assert cli.main(["diff", str(A), str(A), "-q"]) == 0
    assert cli.main(["diff", str(A), str(B), "-q", "--fail-on", "none"]) == 0
    assert cli.main(["diff", str(A), str(B), "-q", "--fail-on", "modified"]) == 1
    # only additions in this pair of edits would not fail --fail-on removed on identical input
    assert cli.main(["diff", str(A), str(A), "-q", "--fail-on", "removed"]) == 0
    assert cli.main(["diff", str(A), str(tmp_path / "nope"), "-q"]) == 2
    assert cli.main(["diff", str(A), str(B), "-q", "--fail-on", "bogus"]) == 2
    assert "ddc: error" in capsys.readouterr().err


def test_cli_no_glb_warns(tmp_path, capsys):
    rc = cli.main(["diff", str(A), str(B), "-q", "--glb", str(tmp_path / "x.glb")])
    assert rc == 1 and not (tmp_path / "x.glb").exists()
    assert "no GLB found" in capsys.readouterr().err


# ---------------------------------------------------------------------------- RVT output with GLBs

def _param_numbers(obj, new):
    """Set the first stored number of a parameter record."""
    if isinstance(obj, dict):
        if "number" in obj:
            obj["number"] = new
            return True
        return any(_param_numbers(v, new) for v in obj.values())
    if isinstance(obj, list):
        return any(_param_numbers(v, new) for v in obj)
    return False


@pytest.fixture()
def rvt_pair(tmp_path):
    a, b = tmp_path / "A.converted", tmp_path / "B.converted"
    shutil.copytree(DATA / "rvt_converted", a)
    b.mkdir()
    shutil.copy(a / "native_graph.ifc", b / "native_graph.ifc")
    lines_out = []
    for ln in (a / "properties.jsonl").read_text(encoding="utf-8").splitlines():
        d = json.loads(ln)
        owner = str(d.get("owner"))
        if owner == "1001":
            continue                                           # wall 1001 is deleted in B
        if owner == "1002" and d.get("row") == "parameter" and d.get("builtin_label") == "Unconnected Height":
            assert _param_numbers(d["fields"], 7.5)             # wall 1002 gets a new height
        lines_out.append(json.dumps(d))
        if owner == "1002":                                    # wall 9000 is new: a copy of 1002
            e = json.loads(json.dumps(d).replace('"1002"', '"9000"'))
            e["owner"] = "9000"
            lines_out.append(json.dumps(e))
    (b / "properties.jsonl").write_text("\n".join(lines_out) + "\n", encoding="utf-8")
    write_cubes(a / "native_graph.glb", [(1001, 10, 0), (1002, 10, 20)])
    write_cubes(b / "native_graph.glb", [(1002, 10, 20), (9000, 10, 40)])
    return a, b


def test_rvt_diff_and_glb_coloring(rvt_pair, tmp_path):
    a, b = rvt_pair
    r = diff.diff_tables(table.build_table(str(a)), table.build_table(str(b)))
    st = {e["key"]: e for e in r["elements"]}
    assert r["key"]["column"] == "GlobalId" or r["key"]["column"] == "ElementId"
    assert {k: v["status"] for k, v in st.items() if k in ("1001", "1002", "9000")} ==         {"1001": "removed", "1002": "modified", "9000": "added"}
    assert [c["field"] for c in st["1002"]["changes"]] == ["Unconnected Height"]
    assert st["1"]["status"] == st["501"]["status"] == "unchanged"
    out = tmp_path / "diff.glb"
    stats = glbdiff.colorize(str(b / "native_graph.glb"), str(a / "native_graph.glb"), r, str(out))
    assert (stats["base_version"], stats["linked"], stats["unlinked"], stats["removed"]) == ("B", 2, 0, 1)
    gltf, blob = glbdiff.read_glb(out)
    status = {n["extras"]["element_id"]: n["extras"]["ddc_diff_status"] for n in gltf["nodes"]}
    assert status == {1002: "modified", 9000: "added", 1001: "removed"}
    colors = {}
    for n in gltf["nodes"]:
        m = gltf["materials"][gltf["meshes"][n["mesh"]]["primitives"][0]["material"]]
        colors[n["extras"]["element_id"]] = m["pbrMetallicRoughness"]["baseColorFactor"][:3]
    assert colors[1001][0] > 0.8 and colors[1001][1] < 0.3                      # red
    assert colors[9000][1] > colors[9000][0] and colors[9000][1] > colors[9000][2]   # green
    assert colors[1002][0] > 0.9 and 0.5 < colors[1002][1] < 0.7                # orange
    # the geometry copied from A reads back through the table reader: a 10 ft cube in metres
    q = table.read_glb_quantities(out)
    assert abs(q["1001"]["volume"]) == pytest.approx((10 * 0.3048) ** 3, rel=1e-5)
    assert abs(q["1002"]["volume"]) == pytest.approx((10 * 0.3048) ** 3, rel=1e-5)
    assert gltf["buffers"][0]["byteLength"] == len(blob)


def test_glb_unchanged_modes_and_cli(tmp_path):
    ta = table.build_table(str(A))
    r = diff.diff_tables(ta, ta)                      # nothing changed: every node is "unchanged"
    g = tmp_path / "in.glb"
    write_cubes(g, [(4131, 3, 0)], extras_key="element_id")
    for mode in ("ghost", "keep", "hide"):
        out = tmp_path / f"{mode}.glb"
        stats = glbdiff.colorize(str(g), None, r, str(out), unchanged=mode)
        gltf, _ = glbdiff.read_glb(out)
        assert stats["unchanged"] == 1
        if mode == "hide":
            assert "mesh" not in gltf["nodes"][0]
            continue
        m = gltf["materials"][gltf["meshes"][gltf["nodes"][0]["mesh"]]["primitives"][0]["material"]]
        alpha = m["pbrMetallicRoughness"]["baseColorFactor"][3]
        assert (alpha < 1 and m.get("alphaMode") == "BLEND") if mode == "ghost" else (alpha == 1 and "alphaMode" not in m)


def test_glb_cli_end_to_end(rvt_pair, tmp_path):
    a, b = rvt_pair
    out = tmp_path / "cli.glb"
    assert cli.main(["diff", str(a), str(b), "-q", "--glb", str(out), "--glb-unchanged", "hide"]) == 1
    gltf, _ = glbdiff.read_glb(out)
    assert {n["extras"]["ddc_diff_status"] for n in gltf["nodes"] if "mesh" in n} == {"modified", "added", "removed"}
    assert gltf["asset"]["extras"]["ddc_diff"]["summary"]["removed"] == 1


def test_glb_links_by_global_id_and_nested_nodes(tmp_path):
    ta, tb = table.build_table(str(A)), table.build_table(str(B))
    r = diff.diff_tables(ta, tb)
    g = tmp_path / "ifc.glb"
    write_cubes(g, [(WALL, 3, 0), (NEW_WALL, 3, 10), (ROOF, 3, 20)], extras_key="global_id", name_prefix="IfcProduct_")
    out = tmp_path / "o.glb"
    stats = glbdiff.colorize(str(g), None, r, str(out))
    assert stats["linked"] == 3 and stats["unlinked"] == 0
    gltf, _ = glbdiff.read_glb(out)
    got = {n["extras"]["global_id"]: n["extras"]["ddc_diff_status"] for n in gltf["nodes"]}
    assert got == {WALL: "modified", NEW_WALL: "added", ROOF: "unchanged"}
    assert stats["removed_missing_geometry"] == 2        # removed elements need a GLB of version A


def test_glb_errors(tmp_path):
    bad = tmp_path / "bad.glb"
    bad.write_bytes(b"not a glb at all, really")
    with pytest.raises(glbdiff.GlbError):
        glbdiff.read_glb(bad)
    assert cli.main(["diff", str(A), str(B), "-q", "--glb", str(tmp_path / "o.glb"), "--glb-b", str(bad)]) == 2
