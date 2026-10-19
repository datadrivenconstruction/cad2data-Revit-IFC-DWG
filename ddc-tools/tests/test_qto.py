import json

from openpyxl import load_workbook

from ddc_tools import qto
from ddc_tools.cli import main


def test_sums_and_counts(rows):
    res = qto.build_qto(rows)
    assert res["quantity_columns"] == ["Length", "Width", "Area", "Volume"]
    walls = next(r for r in res["by_category"] if r["category"] == "Walls")
    assert walls["count"] == 9
    # 2.8 + 1.96 + 6.72 per level for the first six, plus 2.24 + 0 + 0.84
    assert abs(walls["Volume"] - (2 * (2.8 + 1.96 + 6.72) + 2.24 + 0.84)) < 1e-9
    assert sum(r["count"] for r in res["by_level"]) == 18
    assert sum(r["count"] for r in res["detail"]) == 18


def test_levels_and_none_group(rows):
    res = qto.build_qto(rows)
    levels = {r["level"] for r in res["by_level"]}
    assert levels == {"Level 1", "Level 2", "(none)"}


def test_warnings(rows):
    res = qto.build_qto(rows)
    text = " | ".join(f"{w['scope']}: {w['message']}" for w in res["warnings"])
    assert "non-numeric" in text            # "n/a" volume on a column
    assert "empty level" in text
    assert "Generic Models: category has 1 element" in text
    assert "zero" in text                   # wall with zero volume
    ids = next(w for w in res["warnings"] if "non-numeric" in w["message"])["sample_element_ids"]
    assert len(ids) == 1


def test_category_filter_and_explicit_quantities(rows):
    res = qto.build_qto(rows, quantities=["volume", "Nope"], category_filter=["wall*"])
    assert res["quantity_columns"] == ["Volume"] and res["elements"] == 9


def test_no_quantity_columns_still_counts():
    res = qto.build_qto([{"Category": "A"}, {"Category": "A"}])
    assert res["by_category"][0]["count"] == 2
    assert any("no quantity columns" in w["message"] for w in res["warnings"])


def test_empty_input_is_error():
    res = qto.build_qto([])
    assert any(w["level"] == "error" for w in res["warnings"])


def test_ifc_style_quantity_columns():
    rows = [{"IfcClass": "IfcWall", "BuildingStorey": "L1", "Qto_WallBaseQuantities.NetVolume": "2,5"},
            {"IfcClass": "IfcWall", "BuildingStorey": "L1", "Qto_WallBaseQuantities.NetVolume": 1.5}]
    res = qto.build_qto(rows)
    assert res["quantity_columns"] == ["Qto_WallBaseQuantities.NetVolume"]
    assert res["by_category"][0]["Qto_WallBaseQuantities.NetVolume"] == 4.0
    assert res["by_level"][0]["level"] == "L1"


def test_xlsx_layout(export_dir, tmp_path):
    out = tmp_path / "q.xlsx"
    rc = main(["qto", str(export_dir), "--format", "xlsx", "--xlsx", str(out), "-q"])
    assert rc == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Read me", "By Category", "By Type", "By Level", "Detail", "Warnings"]
    ws = wb["By Category"]
    assert [c.value for c in ws[1]][:3] == ["Category", "Count", "Length"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    last = ws.max_row
    assert str(ws.cell(last, 2).value).startswith("=SUBTOTAL(109,")


def test_cli_json_and_strict(export_dir, tmp_path):
    j = tmp_path / "q.json"
    assert main(["qto", str(export_dir), "--json", str(j), "-q"]) == 0
    assert json.loads(j.read_text(encoding="utf-8"))["warnings"]
    assert main(["qto", str(export_dir), "--strict", "-q"]) == 1
    assert main(["qto", str(tmp_path / "nowhere"), "-q"]) == 2
