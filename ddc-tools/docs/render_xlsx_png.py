"""Render the top-left part of an XLSX sheet to a PNG (cell text, fills, widths) for the README.
Dev helper only (needs Pillow); it is a drawing of the cell contents, not an Excel screenshot.

    python render_xlsx_png.py report.xlsx Summary out.png [max_rows] [max_cols]
"""
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


def font(size, bold=False):
    for n in (("segoeuib.ttf" if bold else "segoeui.ttf"), ("arialbd.ttf" if bold else "arial.ttf"),
              ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")):
        try:
            return ImageFont.truetype(n, size)
        except OSError:
            pass
    return ImageFont.load_default()


def render(xlsx, sheet, out, max_rows=22, max_cols=9):
    ws = load_workbook(xlsx)[sheet]
    rows, cols = min(ws.max_row, max_rows), min(ws.max_column, max_cols)
    widths = []
    for c in range(1, cols + 1):
        w = ws.column_dimensions[get_column_letter(c)].width or 10
        widths.append(int(min(w, 46) * 7.2) + 8)
    rh, x0, y0 = 24, 36, 26
    W, H = x0 + sum(widths) + 1, y0 + rows * rh + 1
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    f, fb = font(13), font(13, True)
    d.rectangle([0, 0, W, y0], fill="#F2F2F2")
    x = x0
    for c in range(cols):
        d.text((x + widths[c] // 2 - 4, 5), get_column_letter(c + 1), fill="#444", font=f)
        x += widths[c]
    for r in range(1, rows + 1):
        y = y0 + (r - 1) * rh
        d.rectangle([0, y, x0, y + rh], fill="#F2F2F2", outline="#D0D0D0")
        d.text((8, y + 4), str(r), fill="#444", font=f)
        x = x0
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
            col = "#" + fill[-6:] if isinstance(fill, str) and len(fill) >= 6 else "white"
            d.rectangle([x, y, x + widths[c - 1], y + rh], fill=col, outline="#D9D9D9")
            if cell.value is not None:
                txt = str(cell.value)
                while len(txt) > 1 and d.textlength(txt, font=f) > widths[c - 1] - 8:
                    txt = txt[:-2] + "…" if not txt.endswith("…") else txt[:-2] + "…"
                fc = "white" if col == "#1F3864" else "black"
                d.text((x + 5, y + 4), txt, fill=fc, font=fb if cell.font and cell.font.bold else f)
            x += widths[c - 1]
    im.save(out)


if __name__ == "__main__":
    a = sys.argv
    render(a[1], a[2], a[3], int(a[4]) if len(a) > 4 else 22, int(a[5]) if len(a) > 5 else 9)
