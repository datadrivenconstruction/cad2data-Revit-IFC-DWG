"""ddc check: YAML rules over a converter export -> traffic-light XLSX, JSON, exit code."""
from __future__ import annotations

import fnmatch
import re
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

import yaml
from openpyxl.styles import Font

from . import columns as C
from .paths import fs
from .loader import is_empty
from .xlsxutil import (FILLS, new_workbook, style_header, write_table, autosize, title_cell)

CHECKS = ("required", "not_empty", "pattern", "unique", "range", "enum")
SEVERITIES = ("error", "warning", "info")
MAX_FINDINGS_PER_RULE = 5000


class RuleError(Exception):
    pass


# ---------------------------------------------------------------- rules file

def load_rules(path: str) -> Dict[str, Any]:
    try:
        with open(fs(path), encoding="utf-8") as f:
            spec = yaml.safe_load(f)
    except (OSError, yaml.YAMLError) as e:
        raise RuleError(f"cannot read rules file {path}: {e}")
    return validate_spec(spec)


def _as_list(v: Any) -> List[Any]:
    if v is None:
        return []
    return list(v) if isinstance(v, (list, tuple)) else [v]


def validate_spec(spec: Any) -> Dict[str, Any]:
    if not isinstance(spec, dict) or not isinstance(spec.get("rules"), list) or not spec["rules"]:
        raise RuleError("rules file must be a mapping with a non-empty 'rules' list")
    seen = set()
    for i, r in enumerate(spec["rules"], 1):
        if not isinstance(r, dict):
            raise RuleError(f"rule #{i}: must be a mapping")
        r.setdefault("id", f"R{i:03d}")
        r["id"] = str(r["id"])
        if r["id"] in seen:
            raise RuleError(f"duplicate rule id {r['id']}")
        seen.add(r["id"])
        chk = r.get("check")
        if chk not in CHECKS:
            raise RuleError(f"rule {r['id']}: 'check' must be one of {', '.join(CHECKS)}")
        sev = str(r.get("severity", "error")).lower()
        if sev not in SEVERITIES:
            raise RuleError(f"rule {r['id']}: severity must be one of {', '.join(SEVERITIES)}")
        r["severity"] = sev
        params = _as_list(r.get("params", r.get("param")))
        if not params:
            raise RuleError(f"rule {r['id']}: 'param' or 'params' is required")
        r["params"] = [str(p) for p in params]
        if chk == "pattern":
            if "regex" not in r:
                raise RuleError(f"rule {r['id']}: pattern needs 'regex'")
            try:
                re.compile(str(r["regex"]))
            except re.error as e:
                raise RuleError(f"rule {r['id']}: bad regex: {e}")
        if chk == "range" and "min" not in r and "max" not in r:
            raise RuleError(f"rule {r['id']}: range needs 'min' and/or 'max'")
        if chk == "enum" and not _as_list(r.get("values")):
            raise RuleError(f"rule {r['id']}: enum needs 'values'")
        w = r.get("where") or {}
        if not isinstance(w, dict):
            raise RuleError(f"rule {r['id']}: 'where' must be a mapping")
        for k in w:
            if k not in ("category", "type", "level", "category_regex", "type_regex", "level_regex"):
                raise RuleError(f"rule {r['id']}: unknown 'where' key {k!r}")
            for rx in (_as_list(w[k]) if k.endswith("_regex") else []):
                try:
                    re.compile(str(rx))
                except re.error as e:
                    raise RuleError(f"rule {r['id']}: bad regex in where.{k}: {e}")
        r["where"] = w
    spec.setdefault("name", "DDC rules")
    spec["columns"] = spec.get("columns") or {}
    return spec


# ---------------------------------------------------------------- evaluation

def _glob_match(value: Any, patterns: List[Any]) -> bool:
    if is_empty(value):
        return False
    v = str(value).lower()
    return any(fnmatch.fnmatchcase(v, str(p).lower()) for p in patterns)


def _re_match(value: Any, patterns: List[Any]) -> bool:
    if is_empty(value):
        return False
    return any(re.search(str(p), str(value), re.I) for p in patterns)


def _find_col(columns: List[str], name: str) -> Optional[str]:
    n = C._norm(name)
    for c in columns:
        if c == name:
            return c
    for c in columns:
        if C._norm(c) == n:
            return c
    suff = [c for c in columns if "." in c and C._norm(c.rsplit(".", 1)[1]) == n]
    return suff[0] if len(suff) == 1 else None


def _find_cols(columns: List[str], name: str) -> List[str]:
    """Exact column, else all columns whose name matches ('Pset_X.FireRating' for 'FireRating').
    Several matches are read as one parameter: the first non-empty value counts."""
    one = _find_col(columns, name)
    if one is not None:
        return [one]
    n = C._norm(name)
    return [c for c in columns if "." in c and C._norm(c.rsplit(".", 1)[1]) == n]


def _first(row: Dict[str, Any], cols: List[str]) -> Any:
    for c in cols:
        v = row.get(c)
        if not is_empty(v):
            return v
    return None


def _fmt(v: Any) -> str:
    return "" if is_empty(v) else str(v)


def run_checks(rows: List[Dict[str, Any]], spec: Dict[str, Any], source: str = "",
               files: Optional[List[str]] = None, notes: Optional[List[str]] = None) -> Dict[str, Any]:
    cols = C.all_columns(rows)
    ov = spec.get("columns", {})
    c_id = C.resolve(cols, "id", ov.get("id"))
    c_guid = C.resolve(cols, "guid", ov.get("guid"))
    c_cat = C.resolve(cols, "category", ov.get("category"))
    c_typ = C.resolve(cols, "type", ov.get("type"))
    c_lvl = C.resolve(cols, "level", ov.get("level"))
    roles = {"id": c_id, "guid": c_guid, "category": c_cat, "type": c_typ, "level": c_lvl}

    findings: List[Dict[str, Any]] = []
    rule_results: List[Dict[str, Any]] = []

    for rule in spec["rules"]:
        w = rule["where"]
        applicable = []
        for row in rows:
            if "category" in w and not _glob_match(row.get(c_cat) if c_cat else None, _as_list(w["category"])):
                continue
            if "type" in w and not _glob_match(row.get(c_typ) if c_typ else None, _as_list(w["type"])):
                continue
            if "level" in w and not _glob_match(row.get(c_lvl) if c_lvl else None, _as_list(w["level"])):
                continue
            if "category_regex" in w and not _re_match(row.get(c_cat) if c_cat else None, _as_list(w["category_regex"])):
                continue
            if "type_regex" in w and not _re_match(row.get(c_typ) if c_typ else None, _as_list(w["type_regex"])):
                continue
            if "level_regex" in w and not _re_match(row.get(c_lvl) if c_lvl else None, _as_list(w["level_regex"])):
                continue
            applicable.append(row)

        fails: List[Dict[str, Any]] = []   # (row, param, value, message)
        note = ""
        chk = rule["check"]
        for param in rule["params"]:
            pcols = _find_cols(cols, param)
            if not pcols:
                if chk == "required":
                    for row in applicable:
                        fails.append({"row": row, "param": param, "value": "",
                                      "message": f"parameter '{param}' is not present in the export"})
                    note = (note + f" column '{param}' not found;").strip()
                else:
                    note = (note + f" column '{param}' not found, check skipped;").strip()
                continue
            if chk in ("required", "not_empty"):
                for row in applicable:
                    if is_empty(_first(row, pcols)):
                        fails.append({"row": row, "param": param, "value": "",
                                      "message": f"'{param}' is empty"})
            elif chk == "pattern":
                rx = re.compile(str(rule["regex"]))
                for row in applicable:
                    v = _first(row, pcols)
                    if is_empty(v):
                        continue
                    if not rx.fullmatch(str(v)):
                        fails.append({"row": row, "param": param, "value": _fmt(v),
                                      "message": f"'{param}' = '{v}' does not match {rule['regex']}"})
            elif chk == "enum":
                allowed = {str(x).lower() for x in _as_list(rule["values"])}
                for row in applicable:
                    v = _first(row, pcols)
                    if is_empty(v):
                        continue
                    if str(v).lower() not in allowed:
                        fails.append({"row": row, "param": param, "value": _fmt(v),
                                      "message": f"'{param}' = '{v}' is not one of the allowed values"})
            elif chk == "range":
                lo, hi = rule.get("min"), rule.get("max")
                for row in applicable:
                    v = _first(row, pcols)
                    if is_empty(v):
                        continue
                    n = C.to_number(v)
                    if n is None:
                        fails.append({"row": row, "param": param, "value": _fmt(v),
                                      "message": f"'{param}' = '{v}' is not a number"})
                    elif (lo is not None and n < float(lo)) or (hi is not None and n > float(hi)):
                        rng = f"{'' if lo is None else lo}..{'' if hi is None else hi}"
                        fails.append({"row": row, "param": param, "value": _fmt(v),
                                      "message": f"'{param}' = {v} is outside {rng}"})
            elif chk == "unique":
                scope_col = _find_col(cols, rule["scope"]) if rule.get("scope") else None
                groups: Dict[Any, List[Dict[str, Any]]] = defaultdict(list)
                for row in applicable:
                    v = _first(row, pcols)
                    if is_empty(v):
                        continue
                    key = (str(v), str(row.get(scope_col)) if scope_col else "")
                    groups[key].append(row)
                for (val, _), members in groups.items():
                    if len(members) > 1:
                        for row in members:
                            fails.append({"row": row, "param": param, "value": val,
                                          "message": f"'{param}' = '{val}' is used by {len(members)} elements"})

        failed_ids = {id(f["row"]) for f in fails}
        n_app = len(applicable)
        n_fail = len(failed_ids)
        if n_app == 0:
            status = "SKIP"
            note = (note + " no elements matched 'where'").strip()
        elif not fails and note and chk != "required" and all(
                not _find_cols(cols, p) for p in rule["params"]):
            status = "SKIP"
        elif not fails:
            status = "PASS"
        else:
            status = {"error": "FAIL", "warning": "WARN", "info": "INFO"}[rule["severity"]]

        stored = 0
        for f in fails:
            if stored >= MAX_FINDINGS_PER_RULE:
                break
            row = f["row"]
            findings.append({
                "rule_id": rule["id"], "severity": rule["severity"],
                "element_id": _fmt(row.get(c_id)) if c_id else "",
                "global_id": _fmt(row.get(c_guid)) if c_guid else "",
                "category": _fmt(row.get(c_cat)) if c_cat else "",
                "type": _fmt(row.get(c_typ)) if c_typ else "",
                "level": _fmt(row.get(c_lvl)) if c_lvl else "",
                "parameter": f["param"], "value": f["value"], "message": f["message"],
            })
            stored += 1

        rule_results.append({
            "id": rule["id"], "title": rule.get("title") or rule.get("description") or rule["id"],
            "check": chk, "severity": rule["severity"], "status": status,
            "applicable": n_app, "failed_elements": n_fail, "findings": len(fails),
            "findings_listed": stored,
            "pass_rate": None if n_app == 0 else round(100.0 * (n_app - n_fail) / n_app, 1),
            "note": note,
        })

    statuses = [r["status"] for r in rule_results]
    overall = "FAIL" if "FAIL" in statuses else ("WARN" if "WARN" in statuses else "PASS")
    summary = {
        "overall": overall,
        "rules": len(rule_results),
        "pass": statuses.count("PASS"), "fail": statuses.count("FAIL"),
        "warn": statuses.count("WARN"), "info": statuses.count("INFO"),
        "skipped": statuses.count("SKIP"),
        "elements": len(rows),
        "findings_total": sum(r["findings"] for r in rule_results),
    }
    return {
        "tool": "ddc check", "generated": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "ruleset": spec.get("name"), "source": source, "source_files": files or [],
        "columns_used": roles, "notes": list(notes or []), "summary": summary, "rules": rule_results, "findings": findings,
    }


def exit_code(result: Dict[str, Any], fail_on: str = "error") -> int:
    s = result["summary"]
    if s["fail"]:
        return 1
    if fail_on == "warning" and s["warn"]:
        return 1
    return 0


# ---------------------------------------------------------------- XLSX report

_STATUS_LIGHT = {"PASS": ("PASS", "green"), "FAIL": ("FAIL", "red"), "WARN": ("WARN", "yellow"),
                 "INFO": ("INFO", "blue"), "SKIP": ("SKIP", "grey")}


def write_xlsx(result: Dict[str, Any], path: str) -> None:
    wb, ws = new_workbook("Summary")
    s = result["summary"]
    title_cell(ws, "A1", f"Data check: {result['ruleset']}")
    meta = [("Source", result["source"]), ("Generated", result["generated"]),
            ("Elements", s["elements"]),
            ("Columns used", ", ".join(f"{k}={v}" for k, v in result["columns_used"].items() if v) or "none detected")]
    r = 2
    for k, v in meta:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1
    r += 1
    ws.cell(r, 1, "Overall")
    c = ws.cell(r, 2, s["overall"])
    c.fill = FILLS[_STATUS_LIGHT[s["overall"]][1]]
    c.font = Font(bold=True)
    ws.cell(r, 3, f"{s['pass']} pass, {s['fail']} fail, {s['warn']} warn, {s['info']} info, {s['skipped']} skipped")
    r += 2
    head = ["Status", "Rule", "Title", "Check", "Severity", "Elements checked", "Elements failed",
            "Pass rate %", "Note"]
    for i, h in enumerate(head, 1):
        ws.cell(r, i, h)
    style_header(ws, r, len(head))
    first = r + 1
    for rr in result["rules"]:
        r += 1
        vals = [rr["status"], rr["id"], rr["title"], rr["check"], rr["severity"], rr["applicable"],
                rr["failed_elements"], rr["pass_rate"], rr["note"]]
        for i, v in enumerate(vals, 1):
            ws.cell(r, i, v)
        sc = ws.cell(r, 1)
        sc.fill = FILLS[_STATUS_LIGHT[rr["status"]][1]]
        sc.font = Font(bold=True)
    ws.freeze_panes = ws.cell(first, 1)
    ws.auto_filter.ref = f"A{first - 1}:{ws.cell(r, len(head)).coordinate}"
    ws.cell(r + 2, 1, "Legend: PASS = no failures; FAIL = failures of severity error; "
                      "WARN = failures of severity warning; INFO = notes; SKIP = no elements matched or column missing.")
    autosize(ws, maxw=60)

    wd = wb.create_sheet("Details")
    head = ["Status", "Rule", "Severity", "ElementId", "GlobalId", "Category", "Type", "Level",
            "Parameter", "Value", "Message"]
    sev_status = {"error": "FAIL", "warning": "WARN", "info": "INFO"}
    data = [[sev_status[f["severity"]], f["rule_id"], f["severity"], f["element_id"], f["global_id"],
             f["category"], f["type"], f["level"], f["parameter"], f["value"], f["message"]]
            for f in result["findings"]]
    write_table(wd, head, data)
    for i in range(2, len(data) + 2):
        wd.cell(i, 1).fill = FILLS[_STATUS_LIGHT[wd.cell(i, 1).value][1]]
    autosize(wd, maxw=70)
    trunc = [r_ for r_ in result["rules"] if r_["findings"] > r_["findings_listed"]]
    if trunc:
        wd.cell(len(data) + 3, 1, "Listing truncated for: " + ", ".join(
            f"{t['id']} ({t['findings_listed']} of {t['findings']})" for t in trunc))
    wb.save(fs(path))
