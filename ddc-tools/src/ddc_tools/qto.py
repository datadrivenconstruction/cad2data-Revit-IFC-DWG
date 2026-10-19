"""ddc qto: quantity summary by category / type / level with gap warnings."""
from __future__ import annotations

import fnmatch
from collections import Counter, defaultdict
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from . import columns as C
from .loader import is_empty
from .xlsxutil import FILLS, autosize, new_workbook, style_header, title_cell, write_table

NONE = "(none)"
SAMPLE_IDS = 8


def _label(v: Any) -> str:
    return NONE if is_empty(v) else str(v)


def build_qto(rows: List[Dict[str, Any]], source: str = "", files: Optional[List[str]] = None,
              quantities: Optional[Sequence[str]] = None, category_filter: Optional[Sequence[str]] = None,
              col_override: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    cols = C.all_columns(rows)
    ov = col_override or {}
    c_id = C.resolve(cols, "id", ov.get("id"))
    c_cat = C.resolve(cols, "category", ov.get("category"))
    c_typ = C.resolve(cols, "type", ov.get("type"))
    c_lvl = C.resolve(cols, "level", ov.get("level"))

    if quantities:
        qcols = []
        for q in quantities:
            hit = next((c for c in cols if C._norm(c) == C._norm(q)), None)
            if hit is None:
                hit = next((c for c in cols if "." in c and C._norm(c.rsplit(".", 1)[1]) == C._norm(q)), None)
            if hit is not None:
                qcols.append(hit)
    else:
        qcols = C.quantity_columns(cols)

    warnings: List[Dict[str, Any]] = []

    def warn(level: str, scope: str, message: str, ids: Optional[List[str]] = None) -> None:
        warnings.append({"level": level, "scope": scope, "message": message,
                         "sample_element_ids": ids or []})

    if category_filter:
        pats = [p.lower() for p in category_filter]
        rows = [r for r in rows if c_cat and not is_empty(r.get(c_cat)) and any(
            fnmatch.fnmatchcase(str(r[c_cat]).lower(), p) for p in pats)]

    if not rows:
        warn("error", "input", "no elements to summarise (empty export or filter matched nothing)")
    for role, col in (("category", c_cat), ("type", c_typ), ("level", c_lvl)):
        if col is None:
            warn("warning", "columns", f"no '{role}' column found; everything is grouped under {NONE}")
    if not qcols:
        warn("warning", "columns", "no quantity columns (Length/Area/Volume/...) found; only counts are reported")

    # duplicate ids
    if c_id:
        cnt = Counter(str(r.get(c_id)) for r in rows if not is_empty(r.get(c_id)))
        dup = [k for k, n in cnt.items() if n > 1]
        if dup:
            warn("warning", "ids", f"{len(dup)} element id(s) occur more than once (e.g. multi-model folder); "
                                   "counts include every row", dup[:SAMPLE_IDS])
    else:
        warn("info", "columns", "no element id column found; warnings cannot list element ids")

    def eid(r: Dict[str, Any]) -> str:
        return str(r.get(c_id)) if c_id and not is_empty(r.get(c_id)) else ""

    # numeric values per row
    num: List[Dict[str, Optional[float]]] = []
    bad: Dict[str, List[str]] = defaultdict(list)
    for r in rows:
        d = {}
        for q in qcols:
            v = r.get(q)
            n = C.to_number(v)
            d[q] = n
            if n is None and not is_empty(v):
                bad[q].append(eid(r))
        num.append(d)
    # drop quantity columns that have no number anywhere
    keep = [q for q in qcols if any(d[q] is not None for d in num)]
    for q in qcols:
        if q not in keep:
            warn("warning", q, f"column '{q}' has no numeric values and is left out")
    qcols = keep
    for q, ids in bad.items():
        if q in qcols:
            warn("warning", q, f"{len(ids)} non-numeric value(s) in '{q}' were ignored", ids[:SAMPLE_IDS])

    groups: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
    per_cat: Dict[str, Dict[str, Any]] = {}
    nolabel = {"category": [], "type": [], "level": []}
    for r, d in zip(rows, num):
        cat, typ, lvl = (_label(r.get(c_cat) if c_cat else None), _label(r.get(c_typ) if c_typ else None),
                         _label(r.get(c_lvl) if c_lvl else None))
        for role, v, col in (("category", cat, c_cat), ("type", typ, c_typ), ("level", lvl, c_lvl)):
            if v == NONE and col:
                nolabel[role].append(eid(r))
        for key, store in (((cat, typ, lvl), groups), (cat, per_cat)):
            g = store.setdefault(key, {"count": 0, "sum": {q: 0.0 for q in qcols},
                                       "valid": {q: 0 for q in qcols}, "zero": {q: 0 for q in qcols},
                                       "any_qty": 0, "ids_noqty": []})
            g["count"] += 1
            has = False
            for q in qcols:
                n = d[q]
                if n is not None:
                    g["sum"][q] += n
                    g["valid"][q] += 1
                    has = True
                    if n == 0:
                        g["zero"][q] += 1
            if has:
                g["any_qty"] += 1
            elif len(g["ids_noqty"]) < SAMPLE_IDS:
                g["ids_noqty"].append(eid(r))

    for role, ids in nolabel.items():
        if ids:
            warn("warning", role, f"{len(ids)} element(s) have an empty {role}; grouped under {NONE}",
                 ids[:SAMPLE_IDS])

    for cat, g in sorted(per_cat.items()):
        if qcols and g["any_qty"] == 0:
            warn("warning", cat, f"category has {g['count']} element(s) but no quantities at all "
                                 "(count only; maybe annotation or non-physical elements)", g["ids_noqty"])
            continue
        for q in qcols:
            v, n = g["valid"][q], g["count"]
            if 0 < v < n:
                warn("warning", cat, f"'{q}' is missing for {n - v} of {n} element(s)")
            if g["zero"][q]:
                warn("warning", cat, f"'{q}' is zero for {g['zero'][q]} of {n} element(s)")

    def pack(store, keynames):
        out = []
        for k, g in sorted(store.items(), key=lambda kv: tuple(str(x) for x in (kv[0] if isinstance(kv[0], tuple) else (kv[0],)))):
            kk = k if isinstance(k, tuple) else (k,)
            row = dict(zip(keynames, kk))
            row["count"] = g["count"]
            for q in qcols:
                row[q] = round(g["sum"][q], 6)
            row["_missing"] = {q: g["count"] - g["valid"][q] for q in qcols}
            out.append(row)
        return out

    by_level: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for (cat, typ, lvl), g in groups.items():
        t = by_level.setdefault((lvl, cat), {"count": 0, "sum": {q: 0.0 for q in qcols},
                                             "valid": {q: 0 for q in qcols}})
        t["count"] += g["count"]
        for q in qcols:
            t["sum"][q] += g["sum"][q]
            t["valid"][q] += g["valid"][q]
    by_type: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for (cat, typ, lvl), g in groups.items():
        t = by_type.setdefault((cat, typ), {"count": 0, "sum": {q: 0.0 for q in qcols},
                                            "valid": {q: 0 for q in qcols}})
        t["count"] += g["count"]
        for q in qcols:
            t["sum"][q] += g["sum"][q]
            t["valid"][q] += g["valid"][q]

    return {
        "tool": "ddc qto", "generated": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "source": source, "source_files": files or [], "elements": len(rows),
        "columns_used": {"id": c_id, "category": c_cat, "type": c_typ, "level": c_lvl},
        "quantity_columns": qcols,
        "by_category": pack(per_cat, ["category"]),
        "by_type": pack(by_type, ["category", "type"]),
        "by_level": pack(by_level, ["level", "category"]),
        "detail": pack(groups, ["category", "type", "level"]),
        "warnings": warnings,
    }


# ---------------------------------------------------------------- XLSX

def _sheet(wb, name: str, keys: List[str], key_titles: List[str], rows: List[Dict[str, Any]],
           qcols: List[str]) -> None:
    ws = wb.create_sheet(name)
    header = key_titles + ["Count"] + qcols
    data = [[r[k] for k in keys] + [r["count"]] + [r[q] for q in qcols] for r in rows]
    last = write_table(ws, header, data)
    nk = len(keys)
    for i in range(nk + 1, len(header) + 1):
        for rr in range(2, last + 1):
            c = ws.cell(rr, i)
            c.number_format = "#,##0" if i == nk + 1 else "#,##0.000"
            c.alignment = Alignment(horizontal="right")
    if data:
        tr = last + 1
        ws.cell(tr, 1, "Total (visible rows)").font = Font(bold=True)
        for i in range(nk + 1, len(header) + 1):
            L = get_column_letter(i)
            c = ws.cell(tr, i, f"=SUBTOTAL(109,{L}2:{L}{last})")
            c.font = Font(bold=True)
            c.number_format = "#,##0" if i == nk + 1 else "#,##0.000"
            c.fill = FILLS["grey"]
    autosize(ws, maxw=45)


def write_xlsx(res: Dict[str, Any], path: str) -> None:
    wb, ws = new_workbook("Read me")
    q = res["quantity_columns"]
    title_cell(ws, "A1", "Quantity take-off")
    lines = [
        ("Source", res["source"]),
        ("Generated", res["generated"]),
        ("Elements", res["elements"]),
        ("Quantity columns", ", ".join(q) if q else "none found (counts only)"),
        ("Grouping columns", ", ".join(f"{k}={v}" for k, v in res["columns_used"].items() if v) or "none"),
        ("Warnings", f"{len(res['warnings'])} (see the Warnings sheet)"),
        ("", ""),
        ("Sheets", "By Category, By Type, By Level, Detail (category + type + level), Warnings"),
        ("Units", "Values are summed exactly as exported. They are not converted; check the unit system "
                  "of your export before using them in a bill of quantities."),
        ("Accuracy", "Elements with empty or non-numeric quantities are not counted in the sums. "
                     "The Warnings sheet lists such gaps. Not a certified quantity statement."),
        ("Filtering", "Use the filter buttons; the Total row recalculates for visible rows."),
    ]
    for i, (k, v) in enumerate(lines, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100
    _sheet(wb, "By Category", ["category"], ["Category"], res["by_category"], q)
    _sheet(wb, "By Type", ["category", "type"], ["Category", "Type"], res["by_type"], q)
    _sheet(wb, "By Level", ["level", "category"], ["Level", "Category"], res["by_level"], q)
    _sheet(wb, "Detail", ["category", "type", "level"], ["Category", "Type", "Level"], res["detail"], q)
    wsw = wb.create_sheet("Warnings")
    sev_fill = {"error": "red", "warning": "yellow", "info": "blue"}
    data = [[w["level"].upper(), w["scope"], w["message"], ", ".join(w["sample_element_ids"])]
            for w in res["warnings"]]
    if not data:
        data = [["OK", "", "No gaps found in the quantities", ""]]
    write_table(wsw, ["Level", "Scope", "Message", "Sample ElementIds"], data)
    for i, w in enumerate(res["warnings"], 2):
        wsw.cell(i, 1).fill = FILLS[sev_fill[w["level"]]]
    if not res["warnings"]:
        wsw.cell(2, 1).fill = FILLS["green"]
    autosize(wsw, maxw=90)
    wb.save(path)
