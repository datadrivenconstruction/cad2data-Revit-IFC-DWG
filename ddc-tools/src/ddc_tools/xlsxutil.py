"""Small openpyxl helpers shared by the report writers."""
from __future__ import annotations

from typing import Any, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FILLS = {
    "green": PatternFill("solid", fgColor="C6EFCE"),
    "red": PatternFill("solid", fgColor="FFC7CE"),
    "yellow": PatternFill("solid", fgColor="FFEB9C"),
    "blue": PatternFill("solid", fgColor="DDEBF7"),
    "grey": PatternFill("solid", fgColor="E7E6E6"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")


def new_workbook(first_sheet: str):
    wb = Workbook()
    ws = wb.active
    ws.title = first_sheet
    return wb, ws


def title_cell(ws, ref: str, text: str) -> None:
    c = ws[ref]
    c.value = text
    c.font = Font(bold=True, size=14)


def style_header(ws, row: int, ncols: int) -> None:
    for i in range(1, ncols + 1):
        c = ws.cell(row, i)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)


def write_table(ws, header: Sequence[str], rows: List[Sequence[Any]], start_row: int = 1) -> int:
    """Write header + rows, freeze header, add filter. Returns the last row used."""
    for i, h in enumerate(header, 1):
        ws.cell(start_row, i, h)
    style_header(ws, start_row, len(header))
    r = start_row
    for row in rows:
        r += 1
        for i, v in enumerate(row, 1):
            ws.cell(r, i, v)
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(header))}{r}"
    return r


def autosize(ws, minw: int = 8, maxw: int = 50) -> None:
    widths = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 500)):
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minw, min(maxw, w + 2))
