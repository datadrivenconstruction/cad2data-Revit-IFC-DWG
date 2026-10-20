"""Read a converter export (folder or single file) into a list of flat row dicts.

Supported: .xlsx, .csv, .jsonl/.ndjson, .json (list of objects), .sqlite/.db.
Values are kept as read; empty cells become None. The relational outputs of the IFC and RVT
converters are joined into one flat table by `adapters`.
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
from urllib.parse import quote
from pathlib import Path
from typing import Any, Dict, List, NamedTuple, Optional, Tuple

from . import adapters as A
from .paths import fs

Row = Dict[str, Any]
# rows, quantity-column hint (None = detect by name), notes
Extra = Tuple[List[Row], Optional[List[str]], List[str]]

# Preference when a folder holds several formats of the same export.
_ORDER = [".sqlite", ".sqlite3", ".db", ".jsonl", ".ndjson", ".xlsx", ".csv", ".json"]
# Auxiliary files that are not element tables.
_SKIP_NAMES = {"geometry_index", "property_scopes", "summary", "report", "timings", "identity",
               "tables", "columns", "result"}


class LoadError(Exception):
    pass


def is_empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    return v


def _flatten(obj: Any, prefix: str = "") -> Row:
    out: Row = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}{k}"
            if isinstance(v, dict):
                out.update(_flatten(v, key + "."))
            elif isinstance(v, list):
                out[key] = json.dumps(v, ensure_ascii=False)
            else:
                out[key] = _clean(v)
    return out


def _ws_rows(ws) -> List[Row]:
    it = ws.iter_rows(values_only=True)
    header = None
    for first in it:
        if first and any(c is not None for c in first):
            header = [str(c).strip() if c is not None else "" for c in first]
            break
    if header is None:
        return []
    rows = []
    for r in it:
        if r is None or all(c is None for c in r):
            continue
        rows.append({h: _clean(c) for h, c in zip(header, r) if h})
    return rows


def _read_xlsx(path: Path, sheet: Optional[str]) -> Extra:
    from openpyxl import load_workbook

    wb = load_workbook(fs(path), read_only=True, data_only=True)
    try:
        if not sheet and A.is_ifc_tables(wb.sheetnames):
            tabs = {n: _ws_rows(wb[n]) for n in A.IFC_TABLES if n in wb.sheetnames}
            return A.ifc_rows(tabs)
        if sheet:
            if sheet not in wb.sheetnames:
                raise LoadError(f"sheet {sheet!r} not found in {path.name}: {wb.sheetnames}")
            ws = wb[sheet]
        else:
            ws = wb[wb.sheetnames[0]]
        return _ws_rows(ws), None, []
    finally:
        wb.close()


_SBCS = ["cp1251", "cp1253", "cp1252"]


def _ansi_codepage() -> str:
    import locale
    enc = str(locale.getpreferredencoding(False)).lower()
    return enc if enc in _SBCS else ""


def _csv_encoding(path: Path) -> str:
    """UTF-8 (with or without BOM) first; then UTF-16 (Excel "Unicode Text"), then the Windows code pages
    that Excel uses for "CSV (Comma delimited)" in Cyrillic, Greek and Western European locales."""
    with open(fs(path), "rb") as fh:
        raw = fh.read()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "utf-16"
    try:
        raw.decode("utf-8-sig")
        return "utf-8-sig"
    except UnicodeDecodeError:
        pass
    # The system's own ANSI code page first (what Excel wrote the file in), then the others.
    ansi = _ansi_codepage()
    order = ([ansi] if ansi in _SBCS else []) + [e for e in _SBCS if e != ansi]
    words = {"cp1251": re.compile("[\u0400-\u04ff]{3,}"), "cp1253": re.compile("[\u0370-\u03ff]{3,}")}
    for enc in order:
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        # cp1251/cp1253 decode almost any byte: take them (unless they are the system code page) only when the
        # text has real words in that script; accented Latin letters decoded that way are isolated letters.
        if enc in words and enc != ansi and not words[enc].search(text):
            continue
        return enc
    return "utf-8-sig"


def _read_csv(path: Path) -> Extra:
    csv.field_size_limit(1 << 30)
    with open(fs(path), newline="", encoding=_csv_encoding(path)) as f:
        first = f.readline()
        f.seek(0)
        # delimiter from the header line only; quoting stays standard (no dialect guessing)
        delim = max(",;\t", key=lambda d: first.count(d))
        rd = csv.DictReader(f, delimiter=delim)
        rows = [{(k or "").strip(): _clean(v) for k, v in r.items() if k} for r in rd]
    return rows, None, []


def _read_jsonl(path: Path) -> Extra:
    with open(fs(path), encoding="utf-8-sig") as f:
        head = ""
        for line in f:
            if line.strip():
                head = line
                break
        try:
            first = json.loads(head) if head else {}
        except json.JSONDecodeError:
            first = {}
        if isinstance(first, dict) and "table" in first and "record" in first:
            tabs: Dict[str, List[Row]] = {n: [] for n in A.IFC_TABLES}
            f.seek(0)
            for line in f:
                if not line.strip() or not any(f'"table":"{n}"' in line[:60] for n in A.IFC_TABLES):
                    continue
                d = json.loads(line)
                tabs[d["table"]].append(d["record"])
            return A.ifc_rows(tabs)
        if isinstance(first, dict) and str(first.get("schema", "")).startswith("native_stored_properties"):
            f.seek(0)
            rows, notes = A.rvt_properties_rows(f)
            return rows, None, notes
    rows: List[Row] = []
    with open(fs(path), encoding="utf-8-sig") as f:
        for n, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
            except json.JSONDecodeError as e:
                raise LoadError(f"{path.name}:{n}: invalid JSON ({e.msg})")
            if isinstance(obj, dict):
                rows.append(_flatten(obj))
    return rows, None, []


def _read_json(path: Path) -> Extra:
    with open(fs(path), encoding="utf-8-sig") as f:
        data = json.load(f)
    if isinstance(data, dict):
        for k in ("elements", "rows", "data", "items"):
            if isinstance(data.get(k), list):
                data = data[k]
                break
    if not isinstance(data, list):
        raise LoadError(f"{path.name}: expected a list of objects")
    return [_flatten(o) for o in data if isinstance(o, dict)], None, []


def _read_sqlite(path: Path, table: Optional[str]) -> Extra:
    # percent-encode: a "#", "?" or "%" in the file name would otherwise end the URI path early
    con = sqlite3.connect("file:" + quote(Path(fs(path)).as_posix(), safe="/:") + "?mode=ro", uri=True)
    try:
        con.row_factory = sqlite3.Row
        names = [r[0] for r in con.execute(
            "select name from sqlite_master where type in ('table','view') order by name")]
        if not names:
            raise LoadError(f"{path.name}: no tables")
        if not table and "records" in names:
            cols = {r[1] for r in con.execute('pragma table_info("records")')}
            if {"table_name", "record_json"} <= cols:
                tabs: Dict[str, List[Row]] = {}
                for n in A.IFC_TABLES:
                    tabs[n] = [json.loads(r[0]) for r in con.execute(
                        "select record_json from records where table_name=?", (n,))]
                if A.is_ifc_tables([n for n, v in tabs.items() if v]):
                    return A.ifc_rows(tabs)
        if table:
            if table not in names:
                raise LoadError(f"table {table!r} not found in {path.name}: {names}")
            chosen = table
        else:
            pref = [n for n in names if n.lower() in ("elements", "objects", "data", "properties")]
            if pref:
                chosen = pref[0]
            else:  # otherwise the widest table
                chosen = max(names, key=lambda n: len(con.execute(
                    f'pragma table_info("{n}")').fetchall()))
        cur = con.execute(f'select * from "{chosen}"')
        return [{k: _clean(r[k]) for k in r.keys()} for r in cur], None, []
    finally:
        con.close()


def read_file_ex(path: Path, sheet: Optional[str] = None, table: Optional[str] = None) -> Extra:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        return _read_xlsx(path, sheet)
    if ext == ".csv":
        return _read_csv(path)
    if ext in (".jsonl", ".ndjson"):
        return _read_jsonl(path)
    if ext == ".json":
        return _read_json(path)
    if ext in (".sqlite", ".sqlite3", ".db"):
        return _read_sqlite(path, table)
    raise LoadError(f"unsupported file type: {path.name}")


def read_file(path: Path, sheet: Optional[str] = None, table: Optional[str] = None) -> List[Row]:
    return read_file_ex(path, sheet, table)[0]


def pick_files(folder: Path, fmt: Optional[str] = None) -> List[Path]:
    cands = [p for p in folder.iterdir()
             if p.is_file() and p.suffix.lower() in _ORDER
             and p.stem.lower() not in _SKIP_NAMES and "manifest" not in p.stem.lower()
             and not p.name.startswith("~$")]
    if fmt:
        want = "." + fmt.lower().lstrip(".")
        cands = [p for p in cands if p.suffix.lower() == want]
    cands.sort(key=lambda p: (_ORDER.index(p.suffix.lower()), p.name.lower()))
    return cands


class Loaded(NamedTuple):
    rows: List[Row]
    files: List[str]
    quantity_columns: Optional[List[str]]   # set by adapters that know which columns are quantities
    notes: List[str]


def load_ex(path: str, fmt: Optional[str] = None, sheet: Optional[str] = None,
            table: Optional[str] = None) -> Loaded:
    """Like load(), plus adapter hints. For a folder the preferred format present is used
    and all files of that format are concatenated (one file per model)."""
    p = Path(path)
    if not p.exists():
        raise LoadError(f"path not found: {path}")
    if p.is_file():
        rows, q, notes = read_file_ex(p, sheet, table)
        return Loaded(rows, [p.name], q, notes)
    files = pick_files(p, fmt)
    if not files:
        raise LoadError(f"no export files (xlsx/csv/jsonl/sqlite/json) in {path}")
    ext = files[0].suffix.lower()
    group = [f for f in files if f.suffix.lower() == ext]
    rows: List[Row] = []
    qty: Optional[List[str]] = None
    notes: List[str] = []
    for f in group:
        r_, q, n_ = read_file_ex(f, sheet, table)
        notes.extend(k for k in n_ if k not in notes)
        if q is not None:
            qty = list(dict.fromkeys((qty or []) + q))
        for r in r_:
            if len(group) > 1:
                r.setdefault("_source_file", f.name)
            rows.append(r)
    return Loaded(rows, [f.name for f in group], qty, notes)


def load(path: str, fmt: Optional[str] = None, sheet: Optional[str] = None,
         table: Optional[str] = None) -> Tuple[List[Row], List[str]]:
    r = load_ex(path, fmt, sheet, table)
    return r.rows, r.files
