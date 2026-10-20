"""IFC-like product model with two readers.

* ``from_ifc_excel``: the tables written by DDC IFC to Excel (``data.jsonl``, ``data.sqlite`` or
  ``drawing.xlsx``; sheets/tables Objects, Properties, Quantities, SpatialLinks, TypeLinks, ...).
* ``from_step``: a small reader for the IFC STEP text (``*.ifc``, also the ``native_graph.ifc`` that
  ``RVTconverter --ifc`` writes). It reads only what a flat element table needs: products, property sets,
  quantity sets, spatial containment, type and material links.

Both give the same ``Model`` so the table builder does not care where the data came from.
"""
from __future__ import annotations

import json
import re
import sqlite3
from urllib.parse import quote
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple

from .loader import LoadError
from .paths import fs


@dataclass
class Obj:
    id: str
    cls: str                      # e.g. IfcWall
    guid: Optional[str] = None
    name: Optional[str] = None
    description: Optional[str] = None
    object_type: Optional[str] = None
    long_name: Optional[str] = None
    tag: Optional[str] = None
    is_type: bool = False


@dataclass
class Prop:
    set_name: str
    name: str
    value: Any
    ifc_type: Optional[str] = None


@dataclass
class Qty:
    set_name: str
    name: str
    kind: str                     # Length, Area, Volume, Count, Weight, Time
    value: float


@dataclass
class Model:
    schema: str = ""
    source_kind: str = ""
    objs: Dict[str, Obj] = field(default_factory=dict)
    props: Dict[str, List[Prop]] = field(default_factory=dict)
    quants: Dict[str, List[Qty]] = field(default_factory=dict)
    contained_in: Dict[str, str] = field(default_factory=dict)   # element -> storey/space/...
    aggregated_in: Dict[str, str] = field(default_factory=dict)  # part -> parent
    type_of: Dict[str, str] = field(default_factory=dict)        # occurrence -> type object
    materials: Dict[str, List[str]] = field(default_factory=dict)
    notes: List[str] = field(default_factory=list)


# ----------------------------------------------------------------------------- helpers

SPATIAL = {"IfcProject", "IfcSite", "IfcBuilding", "IfcBuildingStorey"}
SKIP_BY_DEFAULT = SPATIAL | {"IfcOpeningElement"}

_QTY_KIND = {"IfcQuantityLength": "Length", "IfcQuantityArea": "Area", "IfcQuantityVolume": "Volume",
             "IfcQuantityCount": "Count", "IfcQuantityWeight": "Weight", "IfcQuantityTime": "Time"}
_QTY_KEY = {"Length": "LengthValue", "Area": "AreaValue", "Volume": "VolumeValue",
            "Count": "CountValue", "Weight": "WeightValue", "Time": "TimeValue"}


def _s(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return str(v)


def _plain_value(v: Any) -> Tuple[Any, Optional[str]]:
    """Unwrap {'NominalValue': {'ifc_type': 'IfcLabel', 'value': 'x'}} style values."""
    if isinstance(v, dict):
        if "NominalValue" in v:
            return _plain_value(v["NominalValue"])
        if "ifc_type" in v and "value" in v:
            return v["value"], v["ifc_type"]
        for key in ("EnumerationValues", "ListValues"):
            if isinstance(v.get(key), list):
                vals = [_plain_value(x)[0] for x in v[key]]
                return ", ".join(str(x) for x in vals if x is not None), None
        if "value" in v:
            return v["value"], None
        return json.dumps(v, ensure_ascii=False), None
    return v, None


# ----------------------------------------------------------------------------- IFC to Excel tables

_TABLES = ("Objects", "Properties", "Quantities", "SpatialLinks", "TypeLinks", "NormalizedMeasures",
           "Relations", "Entities", "Metadata")
_JSON_COLS = {"value", "path", "Attributes", "attributes", "source_real_values", "defined_types",
              "value_path"}
_MAT_CLASSES = {"IfcMaterial", "IfcMaterialLayer", "IfcMaterialLayerSet", "IfcMaterialLayerSetUsage",
                "IfcMaterialList", "IfcMaterialConstituent", "IfcMaterialConstituentSet"}


def is_ifc_excel(path: Path) -> bool:
    """True if the file/folder holds IFC to Excel tables (not a flat element table)."""
    try:
        for tbl, _ in iter_ifc_excel_records(path, only=("Objects",)):
            return True
    except (LoadError, OSError, sqlite3.Error, ValueError):
        return False
    return False


def _resolve_file(path: Path) -> Path:
    if path.is_file():
        return path
    for name in ("data.jsonl", "data.sqlite", "drawing.xlsx"):
        if (path / name).exists():
            return path / name
    # folder that holds the `<name>-<hash>` sub-folder written by the converter
    subs = sorted(p for p in path.iterdir() if p.is_dir() and (p / "data.jsonl").exists()) if path.is_dir() else []
    if subs:
        return subs[0] / "data.jsonl"
    raise LoadError(f"no IFC to Excel output (data.jsonl / data.sqlite / drawing.xlsx) in {path}")


def iter_ifc_excel_records(path: Path, only: Optional[Iterable[str]] = None) -> Iterator[Tuple[str, dict]]:
    f = _resolve_file(Path(path))
    want = set(only) if only else None
    ext = f.suffix.lower()
    if ext in (".jsonl", ".ndjson"):
        with open(fs(f), encoding="utf-8-sig") as fh:
            for line in fh:
                if not line.strip():
                    continue
                if want is not None:
                    # cheap pre-filter before parsing the (often huge) line
                    head = line[:40]
                    if not any(f'"{t}"' in head for t in want):
                        continue
                d = json.loads(line)
                if "table" in d and "record" in d:
                    if want is None or d["table"] in want:
                        yield d["table"], d["record"]
                else:
                    raise LoadError(f"{f.name}: not an IFC to Excel data.jsonl (no table/record keys)")
    elif ext in (".sqlite", ".sqlite3", ".db"):
        con = sqlite3.connect("file:" + quote(Path(fs(f)).as_posix(), safe="/:") + "?mode=ro", uri=True)
        try:
            names = [r[0] for r in con.execute("select name from sqlite_master where type='table'")]
            if "records" not in names:
                raise LoadError(f"{f.name}: no 'records' table (not IFC to Excel output)")
            q = "select table_name, record_json from records"
            args: tuple = ()
            if want is not None:
                q += " where table_name in (%s)" % ",".join("?" * len(want))
                args = tuple(want)
            q += " order by table_name, ordinal"
            for t, js in con.execute(q, args):
                yield t, json.loads(js)
        finally:
            con.close()
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(f, read_only=True, data_only=True)
        try:
            if "Objects" not in wb.sheetnames:
                raise LoadError(f"{f.name}: no 'Objects' sheet (not IFC to Excel output)")
            for sh in wb.sheetnames:
                if sh not in _TABLES or (want is not None and sh not in want):
                    continue
                it = wb[sh].iter_rows(values_only=True)
                header = next(it, None)
                if not header:
                    continue
                for r in it:
                    rec = {}
                    for h, c in zip(header, r):
                        if h is None:
                            continue
                        if h in _JSON_COLS and isinstance(c, str):
                            try:
                                c = json.loads(c)
                            except ValueError:
                                pass
                        elif isinstance(c, str) and c in ("True", "False"):
                            c = (c == "True")
                        rec[h] = c
                    yield sh, rec
        finally:
            wb.close()
    else:
        raise LoadError(f"unsupported IFC to Excel file: {f.name}")


def _ref_ids(v: Any) -> List[str]:
    out: List[str] = []
    if isinstance(v, dict) and "ref" in v:
        out.append(str(v["ref"]))
    elif isinstance(v, list):
        for x in v:
            out.extend(_ref_ids(x))
    return out


def from_ifc_excel(path: Path) -> Model:
    m = Model(source_kind="ifc-excel")
    norm: Dict[Tuple[str, str], float] = {}
    ent: Dict[str, dict] = {}
    rels: List[dict] = []
    props_rows: List[dict] = []
    for tbl, r in iter_ifc_excel_records(
            path, only=("Objects", "Properties", "Quantities", "SpatialLinks", "TypeLinks",
                        "NormalizedMeasures", "Relations", "Entities", "Metadata")):
        if tbl == "Objects":
            a = r.get("Attributes") or {}
            if isinstance(a, str):
                try:
                    a = json.loads(a)
                except ValueError:
                    a = {}
            sid = str(r["ID"])
            m.objs[sid] = Obj(id=sid, cls=str(r.get("Category") or ""), guid=_s(r.get("GlobalId")),
                              name=_s(r.get("Name")), description=_s(r.get("Description")),
                              object_type=_s(a.get("ObjectType")), long_name=_s(a.get("LongName")),
                              tag=_s(a.get("Tag")), is_type=str(r.get("IsType")).lower() == "true")
        elif tbl == "Properties":
            props_rows.append(r)
        elif tbl == "Quantities":
            kind = _QTY_KIND.get(str(r.get("ifc_class")), "")
            val = r.get("value")
            num = None
            if isinstance(val, dict):
                num = val.get(_QTY_KEY.get(kind, ""))
                if num is None:
                    num = next((x for x in val.values() if isinstance(x, (int, float))), None)
            if kind and isinstance(num, (int, float)):
                m.quants.setdefault(str(r["owner_step_id"]), []).append(
                    Qty(str(r.get("set_name") or ""), str(r.get("property_name") or ""), kind, float(num)))
        elif tbl == "SpatialLinks":
            c, p = str(r["child_step_id"]), str(r["parent_step_id"])
            if r.get("kind") == "IfcRelContainedInSpatialStructure":
                m.contained_in[c] = p
            else:
                m.aggregated_in[c] = p
        elif tbl == "TypeLinks":
            m.type_of[str(r["owner_step_id"])] = str(r["type_step_id"])
        elif tbl == "NormalizedMeasures":
            if r.get("status") == "NORMALIZED_SI" and isinstance(r.get("si_value"), (int, float)):
                norm[(str(r["owner_step_id"]), str(r["property_step_id"]))] = float(r["si_value"])
        elif tbl == "Relations":
            if r.get("ifc_class") == "IfcRelAssociatesMaterial":
                rels.append(r)
        elif tbl == "Entities":
            if r.get("ifc_class") in _MAT_CLASSES:
                ent[str(r["step_id"])] = r
        elif tbl == "Metadata":
            m.schema = str(r.get("ifc_schema") or "")
    for r in props_rows:
        owner = str(r["owner_step_id"])
        val, ity = _plain_value(r.get("value"))
        si = norm.get((owner, str(r.get("property_step_id"))))
        if si is not None and isinstance(val, (int, float)):
            val = si
        m.props.setdefault(owner, []).append(
            Prop(str(r.get("set_name") or ""), str(r.get("property_name") or ""), val, ity))
    # materials
    def mat_names(ref_id: str, depth: int = 0) -> List[str]:
        e = ent.get(ref_id)
        if not e or depth > 4:
            return []
        a = e.get("attributes") or {}
        c = e.get("ifc_class")
        if c == "IfcMaterial":
            return [x for x in [_s(a.get("Name"))] if x]
        if c == "IfcMaterialLayer":
            return [n for i in _ref_ids(a.get("Material")) for n in mat_names(i, depth + 1)]
        if c == "IfcMaterialLayerSet":
            return [n for i in _ref_ids(a.get("MaterialLayers")) for n in mat_names(i, depth + 1)]
        if c == "IfcMaterialLayerSetUsage":
            return [n for i in _ref_ids(a.get("ForLayerSet")) for n in mat_names(i, depth + 1)]
        if c == "IfcMaterialList":
            return [n for i in _ref_ids(a.get("Materials")) for n in mat_names(i, depth + 1)]
        if c == "IfcMaterialConstituent":
            return [n for i in _ref_ids(a.get("Material")) for n in mat_names(i, depth + 1)]
        if c == "IfcMaterialConstituentSet":
            return [n for i in _ref_ids(a.get("MaterialConstituents")) for n in mat_names(i, depth + 1)]
        return []
    for r in rels:
        a = r.get("attributes") or {}
        names = [n for i in _ref_ids(a.get("RelatingMaterial")) for n in mat_names(i)]
        if names:
            for o in _ref_ids(a.get("RelatedObjects")):
                m.materials.setdefault(o, []).extend(names)
    if not m.objs:
        raise LoadError("IFC to Excel output has no Objects table rows")
    return m


# ----------------------------------------------------------------------------- IFC STEP text

_STR_ESC = re.compile(r"\\X2\\([0-9A-Fa-f]+)\\X0\\|\\X\\([0-9A-Fa-f]{2})|\\S\\(.)")
_TOKEN = re.compile(r"""\s*(?:
    (?P<str>'(?:[^']|'')*')
  | (?P<ref>\#\d+)
  | (?P<enum>\.[A-Za-z0-9_]+\.)
  | (?P<num>[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?)
  | (?P<name>[A-Za-z_][A-Za-z0-9_]*)
  | (?P<p>[(),$*=])
)""", re.X)


def _decode_step_string(raw: str) -> str:
    s = raw[1:-1].replace("''", "'")
    if "\\" not in s:
        return s

    def sub(m):
        if m.group(1) is not None:
            h = m.group(1)
            return "".join(chr(int(h[i:i + 4], 16)) for i in range(0, len(h) - 3, 4))
        if m.group(2) is not None:
            return bytes([int(m.group(2), 16)]).decode("latin-1")
        return chr(ord(m.group(3)) + 128)
    return _STR_ESC.sub(sub, s).replace("\\\\", "\\")


class _Typed:
    """A typed value such as IFCLABEL('x')."""
    __slots__ = ("type", "value")

    def __init__(self, type_: str, value: Any):
        self.type = type_
        self.value = value


class _Ref:
    __slots__ = ("id",)

    def __init__(self, i: str):
        self.id = i


def _parse_args(text: str, pos: int) -> Tuple[list, int]:
    """Parse '( ... )' starting at the '(' at ``pos``; returns (list, position after ')')."""
    assert text[pos] == "("
    out: list = []
    pos += 1
    while True:
        m = _TOKEN.match(text, pos)
        if not m:
            raise ValueError(f"bad STEP near {text[pos:pos + 30]!r}")
        pos = m.end()
        if m.group("p") == ")":
            return out, pos
        if m.group("p") == ",":
            continue
        if m.group("str") is not None:
            out.append(_decode_step_string(m.group("str")))
        elif m.group("ref") is not None:
            out.append(_Ref(m.group("ref")[1:]))
        elif m.group("enum") is not None:
            out.append(m.group("enum"))
        elif m.group("num") is not None:
            t = m.group("num")
            out.append(float(t) if any(c in t for c in ".eE") else int(t))
        elif m.group("p") == "$":
            out.append(None)
        elif m.group("p") == "*":
            out.append("*")
        elif m.group("p") == "(":
            sub, pos = _parse_args(text, pos - 1)
            out.append(sub)
        elif m.group("name") is not None:
            nxt = _TOKEN.match(text, pos)
            if nxt and nxt.group("p") == "(":
                sub, pos = _parse_args(text, nxt.end() - 1)
                out.append(_Typed(m.group("name").upper(), sub[0] if len(sub) == 1 else sub))
            else:
                out.append(m.group("name"))
        else:
            raise ValueError("unexpected token")


_GUID = re.compile(r"[0-9A-Za-z_$]{22}")
_STATEMENT = re.compile(r"(?:'(?:[^']|'')*'|[^';])*;")
_ENTITY = re.compile(r"\s*#(\d+)\s*=\s*([A-Za-z0-9_]+)\s*(\(.*\))\s*;\s*$", re.S)


def _iter_entities(text: str) -> Iterator[Tuple[str, str, list]]:
    i = text.find("DATA;")
    if i < 0:
        raise LoadError("not an IFC STEP file (no DATA section)")
    body = text[i + 5:]
    j = body.rfind("ENDSEC;")
    if j >= 0:
        body = body[:j]
    for st in _STATEMENT.finditer(body):
        s = st.group(0)
        m = _ENTITY.match(s)
        if not m:
            continue
        try:
            args, _ = _parse_args(m.group(3), 0)
        except (ValueError, AssertionError):
            continue
        yield m.group(1), m.group(2).upper(), args


def _nz(v: Any) -> Optional[str]:
    return _s(v) if isinstance(v, str) else None


def _refs(v: Any) -> List[str]:
    if isinstance(v, _Ref):
        return [v.id]
    if isinstance(v, list):
        return [r for x in v for r in _refs(x)]
    return []


_STEP_QTY = {"IFCQUANTITYLENGTH": "Length", "IFCQUANTITYAREA": "Area", "IFCQUANTITYVOLUME": "Volume",
             "IFCQUANTITYCOUNT": "Count", "IFCQUANTITYWEIGHT": "Weight", "IFCQUANTITYTIME": "Time"}
# classes with a GlobalId that are not elements of the model
_NOT_PRODUCT_PREFIX = ("IFCREL", "IFCPROPERTY", "IFCQUANTITY", "IFCCOMPLEX", "IFCPHYSICAL", "IFCMATERIAL")
_NOT_PRODUCT = {"IFCPROPERTYSET", "IFCELEMENTQUANTITY", "IFCPROPERTYSETTEMPLATE", "IFCPROPERTYTEMPLATE",
                "IFCDOORLININGPROPERTIES", "IFCWINDOWLININGPROPERTIES", "IFCDOORPANELPROPERTIES",
                "IFCWINDOWPANELPROPERTIES", "IFCSHAPEASPECT", "IFCGROUP", "IFCSYSTEM", "IFCZONE",
                "IFCDOORSTYLE_", "IFCWINDOWSTYLE_"}


def _pretty_class(upper: str, known: Dict[str, str]) -> str:
    return known.get(upper) or ("Ifc" + upper[3:].title() if upper.startswith("IFC") else upper)


def from_step(path: Path, class_names: Optional[Dict[str, str]] = None) -> Model:
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    m = Model(source_kind="ifc-step")
    sm = re.search(r"FILE_SCHEMA\s*\(\s*\(\s*'([^']+)'", text)
    m.schema = sm.group(1) if sm else ""
    ents: Dict[str, Tuple[str, list]] = {}
    for eid, cls, args in _iter_entities(text):
        ents[eid] = (cls, args)
    # real class spelling (IfcWallStandardCase) is not in the file: rebuild CamelCase from a table of
    # the usual names; unknown ones fall back to Title-case which is right for most one-word classes
    names = _class_name_table()
    if class_names:
        names.update(class_names)

    def pretty(c: str) -> str:
        return names.get(c) or _pretty_class(c, names)

    def is_type_class(c: str) -> bool:
        return c.endswith("TYPE") or c.endswith("STYLE") or c == "IFCTYPEOBJECT" or c == "IFCTYPEPRODUCT"

    for eid, (cls, a) in ents.items():
        if cls.startswith(_NOT_PRODUCT_PREFIX) or cls in _NOT_PRODUCT:
            continue
        if len(a) >= 4 and isinstance(a[0], str) and _GUID.fullmatch(a[0]) and isinstance(a[1], (_Ref, type(None))) \
                and cls not in ("IFCOWNERHISTORY",):
            o = Obj(id=eid, cls=pretty(cls), guid=a[0], name=_nz(a[2]), description=_nz(a[3]),
                    is_type=is_type_class(cls))
            if not o.is_type and len(a) > 4:
                o.object_type = _nz(a[4])
            if cls == "IFCSPACE" and len(a) > 7:
                o.long_name = _nz(a[7])
            if not o.is_type and len(a) > 7 and cls not in ("IFCSPACE", "IFCPROJECT", "IFCSITE",
                                                            "IFCBUILDING", "IFCBUILDINGSTOREY"):
                o.tag = _nz(a[7])
            m.objs[eid] = o

    def prop_value(p) -> Tuple[Any, Optional[str]]:
        v = p
        if isinstance(v, _Typed):
            inner = v.value
            if isinstance(inner, str) and v.type in ("IFCBOOLEAN", "IFCLOGICAL"):
                inner = inner
            return inner, "Ifc" + v.type[3:].title()
        if isinstance(v, str) and v.startswith(".") and v.endswith("."):
            return {"T": True, "F": False}.get(v[1:-1], v[1:-1]), None
        return v, None

    def read_prop(pid: str, set_name: str, out: List[Prop]) -> None:
        e = ents.get(pid)
        if not e:
            return
        c, a = e
        if c == "IFCPROPERTYSINGLEVALUE" and len(a) >= 3:
            v, t = prop_value(a[2])
            if isinstance(v, str) and v in (".T.", ".F."):
                v = (v == ".T.")
            out.append(Prop(set_name, a[0] or "", v, t))
        elif c == "IFCPROPERTYENUMERATEDVALUE" and len(a) >= 3 and isinstance(a[2], list):
            vals = [prop_value(x)[0] for x in a[2]]
            out.append(Prop(set_name, a[0] or "", ", ".join(str(x) for x in vals if x is not None)))
        elif c == "IFCPROPERTYLISTVALUE" and len(a) >= 3 and isinstance(a[2], list):
            vals = [prop_value(x)[0] for x in a[2]]
            out.append(Prop(set_name, a[0] or "", ", ".join(str(x) for x in vals if x is not None)))
        elif c == "IFCPROPERTYBOUNDEDVALUE" and len(a) >= 4:
            v, t = prop_value(a[3] if a[3] is not None else a[2])
            out.append(Prop(set_name, a[0] or "", v, t))
        elif c == "IFCCOMPLEXPROPERTY" and len(a) >= 4:
            for sub in _refs(a[3]):
                read_prop(sub, f"{set_name}.{a[0]}", out)

    psets: Dict[str, List[Prop]] = {}
    qsets: Dict[str, List[Qty]] = {}
    for eid, (cls, a) in ents.items():
        if cls == "IFCPROPERTYSET" and len(a) >= 5:
            ps: List[Prop] = []
            for pid in _refs(a[4]):
                read_prop(pid, a[2] or "", ps)
            psets[eid] = ps
        elif cls == "IFCELEMENTQUANTITY" and len(a) >= 6:
            qs: List[Qty] = []
            for qid in _refs(a[5]):
                q = ents.get(qid)
                if q and q[0] in _STEP_QTY and len(q[1]) >= 4 and isinstance(q[1][3], (int, float)):
                    qs.append(Qty(a[2] or "", q[1][0] or "", _STEP_QTY[q[0]], float(q[1][3])))
            qsets[eid] = qs

    def mats(rid: str, depth: int = 0) -> List[str]:
        e = ents.get(rid)
        if not e or depth > 4:
            return []
        c, a = e
        if c == "IFCMATERIAL":
            return [a[0]] if a and a[0] else []
        if c == "IFCMATERIALLAYER":
            return mats(a[0].id, depth + 1) if a and isinstance(a[0], _Ref) else []
        if c == "IFCMATERIALLAYERSET":
            return [n for i in _refs(a[0]) for n in mats(i, depth + 1)]
        if c == "IFCMATERIALLAYERSETUSAGE":
            return [n for i in _refs(a[0]) for n in mats(i, depth + 1)]
        if c == "IFCMATERIALLIST":
            return [n for i in _refs(a[0]) for n in mats(i, depth + 1)]
        if c == "IFCMATERIALCONSTITUENT":
            return mats(a[2].id, depth + 1) if len(a) > 2 and isinstance(a[2], _Ref) else []
        if c == "IFCMATERIALCONSTITUENTSET":
            return [n for i in _refs(a[2] if len(a) > 2 else None) for n in mats(i, depth + 1)]
        return []

    for eid, (cls, a) in ents.items():
        if cls == "IFCRELDEFINESBYPROPERTIES" and len(a) >= 6:
            rel = _refs(a[5])
            for o in _refs(a[4]):
                for d in rel:
                    if d in psets:
                        m.props.setdefault(o, []).extend(psets[d])
                    elif d in qsets:
                        m.quants.setdefault(o, []).extend(qsets[d])
        elif cls == "IFCRELDEFINESBYTYPE" and len(a) >= 6:
            for o in _refs(a[4]):
                for t in _refs(a[5]):
                    m.type_of[o] = t
        elif cls == "IFCRELCONTAINEDINSPATIALSTRUCTURE" and len(a) >= 6:
            for o in _refs(a[4]):
                for p in _refs(a[5]):
                    m.contained_in[o] = p
        elif cls == "IFCRELAGGREGATES" and len(a) >= 6:
            for o in _refs(a[5]):
                for p in _refs(a[4]):
                    m.aggregated_in[o] = p
        elif cls == "IFCRELASSOCIATESMATERIAL" and len(a) >= 6:
            names = [n for i in _refs(a[5]) for n in mats(i)]
            if names:
                for o in _refs(a[4]):
                    m.materials.setdefault(o, []).extend(names)
    # IFC2x3 types carry property sets in HasPropertySets (index 5): attach them to the type object
    for eid, o in m.objs.items():
        if o.is_type:
            a = ents[eid][1]
            if len(a) > 5 and isinstance(a[5], list):
                for d in _refs(a[5]):
                    if d in psets:
                        m.props.setdefault(eid, []).extend(psets[d])
    if not any(not o.is_type for o in m.objs.values()):
        raise LoadError("no products found in the IFC file")
    return m


_CLASS_TABLE: Optional[Dict[str, str]] = None


def _class_name_table() -> Dict[str, str]:
    """UPPERCASE -> IfcCamelCase for the classes people actually filter on."""
    global _CLASS_TABLE
    if _CLASS_TABLE is None:
        names = """IfcProject IfcSite IfcBuilding IfcBuildingStorey IfcSpace IfcWall IfcWallStandardCase IfcWallElementedCase
        IfcSlab IfcRoof IfcBeam IfcColumn IfcMember IfcPlate IfcFooting IfcPile IfcRailing IfcStair IfcStairFlight
        IfcRamp IfcRampFlight IfcDoor IfcWindow IfcCovering IfcCurtainWall IfcOpeningElement IfcBuildingElementProxy
        IfcFurnishingElement IfcFurniture IfcSystemFurnitureElement IfcFlowTerminal IfcFlowSegment IfcFlowFitting
        IfcFlowController IfcFlowMovingDevice IfcFlowStorageDevice IfcFlowTreatmentDevice IfcEnergyConversionDevice
        IfcDistributionElement IfcDistributionFlowElement IfcDistributionControlElement IfcDistributionPort
        IfcDuctSegment IfcDuctFitting IfcPipeSegment IfcPipeFitting IfcCableCarrierSegment IfcCableSegment
        IfcLightFixture IfcSanitaryTerminal IfcElectricAppliance IfcAirTerminal IfcTransportElement IfcGeographicElement
        IfcReinforcingBar IfcReinforcingMesh IfcTendon IfcFastener IfcMechanicalFastener IfcDiscreteAccessory
        IfcAnnotation IfcGrid IfcVirtualElement IfcElementAssembly IfcCivilElement IfcSensor IfcActuator
        IfcWallType IfcSlabType IfcBeamType IfcColumnType IfcMemberType IfcPlateType IfcFootingType IfcRailingType
        IfcStairType IfcStairFlightType IfcRampType IfcRampFlightType IfcDoorType IfcWindowType IfcCoveringType
        IfcDoorStyle IfcWindowStyle IfcFurnitureType IfcFurnishingElementType IfcOpeningElement IfcRoofType
        IfcBuildingElementProxyType IfcFlowTerminalType IfcFlowSegmentType IfcFlowFittingType IfcCurtainWallType
        IfcTransportElementType IfcDistributionElementType IfcElementType IfcSpaceType IfcTypeObject IfcTypeProduct
        IfcLightFixtureType IfcSanitaryTerminalType IfcDuctSegmentType IfcPipeSegmentType IfcAirTerminalType""".split()
        _CLASS_TABLE = {n.upper(): n for n in names}
    return dict(_CLASS_TABLE)


# ----------------------------------------------------------------------------- detection

def sniff_kind(path: Path) -> Optional[str]:
    """'ifc-step' for a STEP text file, 'ifc-excel' for DDC IFC to Excel tables, else None."""
    p = Path(path)
    if p.is_file() and p.suffix.lower() in (".ifc",):
        return "ifc-step"
    try:
        if is_ifc_excel(p):
            return "ifc-excel"
    except LoadError:
        return None
    return None
