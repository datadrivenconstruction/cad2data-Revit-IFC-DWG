"""BCF issue files from ``ddc check`` and ``ddc diff`` findings, and a reader that tabulates any BCF.

BCF (BIM Collaboration Format) is an open exchange format of buildingSMART International. This module
writes BCF-XML 2.1 and 3.0 zip files from the JSON reports of ``ddc check`` and ``ddc diff`` and reads
BCF 2.0 / 2.1 / 3.0 files into a table. Written from the published schemas and specification only. The
files are checked against the official XSD schemas in the tests, but no certification is claimed.

What is written:

* one topic per rule (check) or per status and category (diff), or one per element with ``per_element``;
* a viewpoint with the affected components: ``IfcGuid`` when the element has a valid IFC GlobalId,
  otherwise ``AuthoringToolId`` (the Revit ElementId or the table id) with ``OriginatingSystem``;
  for diff the components are coloured by status;
* no snapshot. A camera is written only when given (BCF 3.0 schema requires one in a viewpoint, so
  3.0 files carry a neutral default camera that a viewer replaces by zooming to the selection).
"""
from __future__ import annotations

import io
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from xml.etree import ElementTree as ET

from .diff import STATUS_COLORS, STATUSES
from .xlsxutil import autosize, new_workbook, title_cell, write_table

VERSIONS = ("2.1", "3.0")
IFC_GUID = re.compile(r"^[0-9A-Za-z_$]{22}$")
ORIGIN = "DDC ddc-tools"
_NS = uuid.UUID("6f1f3c6e-2f0a-5d0e-9b53-0c5a3dd0b2a1")     # namespace for deterministic ids
MAX_ZIP_MEMBER = 200 * 1024 * 1024
MAX_ZIP_MEMBERS = 200000

SEVERITY_PRIORITY = {"error": "High", "warning": "Normal", "info": "Low"}
SEVERITY_COLOR = {"error": "DB3232", "warning": "FA9E1A", "info": "5B9BD5"}
STATUS_TEXT = {"added": "Added in B (new version)", "removed": "Removed in B (present only in A)",
               "modified": "Modified between A and B", "unchanged": "Unchanged"}


class BcfError(Exception):
    pass


def _hex(rgb: Tuple[int, int, int]) -> str:
    return "%02X%02X%02X" % rgb


def _guid(*parts: str) -> str:
    return str(uuid.uuid5(_NS, "|".join(parts)))


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _txt(s: Any, limit: int = 20000) -> str:
    s = "" if s is None else str(s)
    s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)
    return s[:limit]


# ------------------------------------------------------------------------------------------ topics

class Topic:
    def __init__(self, key: str, title: str, description: str = "", priority: str = "Normal",
                 ttype: str = "Issue", labels: Sequence[str] = ()):
        self.key = key
        self.title = title
        self.description = description
        self.priority = priority
        self.type = ttype
        self.labels = list(dict.fromkeys(l for l in labels if l))
        self.components: List[Tuple[str, str]] = []       # (ifc guid or "", authoring tool id or "")
        self.colors: Dict[str, List[Tuple[str, str]]] = {}
        self.truncated = 0

    def add(self, comp: Tuple[str, str], color: Optional[str], limit: int) -> None:
        if comp == ("", ""):
            return
        if len(self.components) >= limit:
            self.truncated += 1
            return
        self.components.append(comp)
        if color:
            self.colors.setdefault(color, []).append(comp)


def _component(guid: Any, authoring: Any) -> Tuple[str, str]:
    g = str(guid or "").strip()
    a = str(authoring or "").strip()
    if IFC_GUID.match(g):
        return g, ""
    return "", a or g


def topics_from_check(report: Dict[str, Any], per_element: bool = False, max_components: int = 1000) -> List[Topic]:
    rules = {r["id"]: r for r in report.get("rules", [])}
    by_rule: Dict[str, List[Dict[str, Any]]] = {}
    for f in report.get("findings", []):
        by_rule.setdefault(f["rule_id"], []).append(f)
    out: List[Topic] = []
    for rid, fs in by_rule.items():
        rule = rules.get(rid, {})
        sev = fs[0].get("severity", rule.get("severity", "error"))
        title = rule.get("title") or rid
        color = SEVERITY_COLOR.get(sev, SEVERITY_COLOR["error"])
        label = [ "ddc-check", rid, sev ]
        if per_element:
            for n, f in enumerate(fs):
                t = Topic(f"check|{rid}|{n}|{f.get('global_id')}|{f.get('element_id')}",
                          _txt(f"[{rid}] {title}: {f.get('category') or 'element'} {f.get('element_id') or f.get('global_id')}", 300),
                          _txt(f"{f.get('message', '')}\nRule {rid} ({rule.get('check', '')}, {sev}); "
                               f"parameter {f.get('parameter')}; value {f.get('value')!r}; level {f.get('level') or '-'}; "
                               f"type {f.get('type') or '-'}"),
                          SEVERITY_PRIORITY.get(sev, "Normal"), "Error" if sev == "error" else "Warning" if sev == "warning" else "Info",
                          label + [f.get("category", "")])
                t.add(_component(f.get("global_id"), f.get("element_id")), color, max_components)
                out.append(t)
            continue
        examples = list(dict.fromkeys(f.get("message", "") for f in fs))[:12]
        cats = sorted({f.get("category") for f in fs if f.get("category")})
        desc = (f"Rule {rid}: {title}\nCheck: {rule.get('check', '')}; severity: {sev}.\n"
                f"{len(fs)} finding(s) on {len({(f.get('global_id') or f.get('element_id')) for f in fs})} element(s)"
                + (f"; categories: {', '.join(cats[:8])}" if cats else "") + ".\n\nExamples:\n- " + "\n- ".join(examples))
        t = Topic(f"check|{rid}", _txt(f"[{rid}] {title}", 300), _txt(desc), SEVERITY_PRIORITY.get(sev, "Normal"),
                  "Error" if sev == "error" else "Warning" if sev == "warning" else "Info", label)
        seen = set()
        for f in fs:
            k = (f.get("global_id"), f.get("element_id"))
            if k in seen:
                continue
            seen.add(k)
            t.add(_component(f.get("global_id"), f.get("element_id")), color, max_components)
        if t.truncated:
            t.description += f"\n\n{t.truncated} more element(s) are not selected in the viewpoint (limit {max_components})."
        out.append(t)
    return out


def _change_line(c: Dict[str, Any]) -> str:
    old = c["old"] if c["old"] != "" else "(empty)"
    new = c["new"] if c["new"] != "" else "(empty)"
    extra = f" ({c['delta']:+g})" if isinstance(c.get("delta"), (int, float)) else ""
    return f"{c['field']}: {old} -> {new}{extra}"


def topics_from_diff(report: Dict[str, Any], per_element: bool = False, max_components: int = 1000,
                     only: Sequence[str] = ("added", "removed", "modified")) -> List[Topic]:
    elems = [e for e in report.get("elements", []) if e["status"] in only]
    out: List[Topic] = []
    prio = {"removed": "High", "modified": "Normal", "added": "Normal", "unchanged": "Low"}
    if per_element:
        for e in elems:
            who = e.get("type") or e.get("name") or e["category"]
            lines = [STATUS_TEXT[e["status"]], f"Category: {e['category']}; type: {e.get('type') or '-'}; "
                     f"level: {e.get('level') or '-'}; key: {e['key']}"]
            lines += ["- " + _change_line(c) for c in e["changes"][:40]]
            t = Topic(f"diff|{e['status']}|{e['key']}|{e.get('global_id')}", _txt(f"{e['status'].capitalize()}: {e['category']} {who}", 300),
                      _txt("\n".join(lines)), prio[e["status"]], "Change", ["ddc-diff", e["status"], e["category"]])
            t.add(_component(e.get("global_id"), e.get("element_id")), _hex(STATUS_COLORS[e["status"]]), max_components)
            out.append(t)
        return out
    groups: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
    for e in elems:
        groups.setdefault((e["status"], e["category"]), []).append(e)
    order = {s: n for n, s in enumerate(STATUSES)}
    for (st, cat), es in sorted(groups.items(), key=lambda kv: (order[kv[0][0]], kv[0][1])):
        lines = [STATUS_TEXT[st], f"Category: {cat}; {len(es)} element(s).", ""]
        for e in es[:15]:
            head = f"{e.get('type') or e.get('name') or ''} [{e['key']}]".strip()
            lines.append(f"* {head}" + (f", level {e['level']}" if e.get("level") else ""))
            for c in e["changes"][:6]:
                lines.append("    - " + _change_line(c))
            if len(e["changes"]) > 6:
                lines.append(f"    - ... {len(e['changes']) - 6} more change(s)")
        if len(es) > 15:
            lines.append(f"... and {len(es) - 15} more element(s).")
        if st == "removed":
            lines.append("\nRemoved elements are not in the new model; they are listed here by their id in the old one.")
        t = Topic(f"diff|{st}|{cat}", _txt(f"{st.capitalize()}: {cat} ({len(es)})", 300), _txt("\n".join(lines)),
                  prio[st], "Change", ["ddc-diff", st, cat])
        for e in es:
            t.add(_component(e.get("global_id"), e.get("element_id")), _hex(STATUS_COLORS[st]), max_components)
        if t.truncated:
            t.description += f"\n\n{t.truncated} more element(s) are not selected in the viewpoint (limit {max_components})."
        out.append(t)
    return out


def topics_from_report(report: Dict[str, Any], per_element: bool = False, max_components: int = 1000,
                       only: Sequence[str] = ("added", "removed", "modified")) -> Tuple[List[Topic], str]:
    tool = report.get("tool")
    if tool == "ddc check":
        return topics_from_check(report, per_element, max_components), "check"
    if tool == "ddc diff":
        return topics_from_diff(report, per_element, max_components, only), "diff"
    raise BcfError("input must be a JSON report written by `ddc check --json` or `ddc diff --json`")


# ------------------------------------------------------------------------------------------ writer

def _sub(parent: ET.Element, tag: str, text: Optional[str] = None, **attrs: str) -> ET.Element:
    e = ET.SubElement(parent, tag, {k: v for k, v in attrs.items() if v is not None})
    if text is not None:
        e.text = text
    return e


def _xml(root: ET.Element) -> bytes:
    ET.indent(root)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _camera(parent: ET.Element, cam: Dict[str, Any], version: str) -> None:
    c = _sub(parent, "PerspectiveCamera")
    for name, v in (("CameraViewPoint", cam["position"]), ("CameraDirection", cam["direction"]),
                    ("CameraUpVector", cam["up"])):
        p = _sub(c, name)
        for axis, val in zip("XYZ", v):
            _sub(p, axis, repr(float(val)))
    _sub(c, "FieldOfView", repr(float(cam.get("fov", 60.0))))
    if version == "3.0":                       # AspectRatio exists from BCF 3.0 on
        _sub(c, "AspectRatio", repr(float(cam.get("aspect", 16 / 9))))


DEFAULT_CAMERA = {"position": (25.0, -25.0, 18.0), "direction": (-0.58, 0.58, -0.57), "up": (0.0, 0.0, 1.0),
                  "fov": 60.0, "aspect": 16 / 9}


def parse_camera(spec: str) -> Dict[str, Any]:
    try:
        v = [float(x) for x in spec.split(",")]
    except ValueError:
        v = []
    if len(v) != 6:
        raise BcfError("camera must be six numbers: x,y,z,dx,dy,dz (position and viewing direction, metres)")
    return {"position": tuple(v[:3]), "direction": tuple(v[3:]), "up": (0.0, 0.0, 1.0), "fov": 60.0, "aspect": 16 / 9}


def _comp(parent: ET.Element, c: Tuple[str, str]) -> None:
    attrs = {"IfcGuid": c[0]} if c[0] else {}
    e = _sub(parent, "Component", **attrs)
    if c[1]:
        _sub(e, "OriginatingSystem", ORIGIN)
        _sub(e, "AuthoringToolId", c[1])


def _viewpoint(t: Topic, vp_guid: str, version: str, camera: Optional[Dict[str, Any]], default_visible: bool = True) -> bytes:
    root = ET.Element("VisualizationInfo", {"Guid": vp_guid})
    comps = _sub(root, "Components")
    sel = _sub(comps, "Selection")
    for c in t.components:
        _comp(sel, c)
    _sub(comps, "Visibility", DefaultVisibility="true" if default_visible else "false")
    if t.colors:
        col = _sub(comps, "Coloring")
        for color, cs in t.colors.items():
            ce = _sub(col, "Color", Color=color)
            holder = _sub(ce, "Components") if version == "3.0" else ce
            for c in cs:
                _comp(holder, c)
    cam = camera or (DEFAULT_CAMERA if version == "3.0" else None)
    if cam:
        _camera(root, cam, version)
    return _xml(root)


def _markup(t: Topic, topic_guid: str, vp_guid: Optional[str], version: str, author: str, now: str,
            file_name: Optional[str]) -> bytes:
    root = ET.Element("Markup")
    if file_name:
        h = _sub(root, "Header")
        parent = _sub(h, "Files") if version == "3.0" else h
        f = _sub(parent, "File", **({"IsExternal": "true"} if version == "3.0" else {"isExternal": "true"}))
        _sub(f, "Filename", file_name)
    topic = _sub(root, "Topic", Guid=topic_guid, TopicType=t.type, TopicStatus="Open")
    _sub(topic, "Title", t.title)
    _sub(topic, "Priority", t.priority)
    if version == "3.0":
        if t.labels:
            ls = _sub(topic, "Labels")
            for l in t.labels:
                _sub(ls, "Label", l)
    else:
        for l in t.labels:
            _sub(topic, "Labels", l)
    _sub(topic, "CreationDate", now)
    _sub(topic, "CreationAuthor", author)
    if t.description:
        _sub(topic, "Description", t.description)
    if version == "3.0":
        if vp_guid:
            vps = _sub(topic, "Viewpoints")
            v = _sub(vps, "ViewPoint", Guid=vp_guid)
            _sub(v, "Viewpoint", "viewpoint.bcfv")
            _sub(v, "Index", "0")
    elif vp_guid:
        v = _sub(root, "Viewpoints", Guid=vp_guid)
        _sub(v, "Viewpoint", "viewpoint.bcfv")
        _sub(v, "Index", "0")
    return _xml(root)


def write_bcf(topics: Sequence[Topic], path: str, version: str = "2.1", author: str = "ddc-tools",
              project_name: str = "DDC findings", file_name: Optional[str] = None,
              camera: Optional[Dict[str, Any]] = None, seed: str = "") -> Dict[str, Any]:
    if version not in VERSIONS:
        raise BcfError(f"BCF version must be one of {', '.join(VERSIONS)}")
    now = _now()
    project_id = _guid("project", seed or project_name)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    n_vp = 0
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        ver = ET.Element("Version", {"VersionId": version})
        if version == "2.1":
            _sub(ver, "DetailedVersion", "2.1")
        z.writestr("bcf.version", _xml(ver))
        if version == "3.0":
            p = ET.Element("ProjectInfo")
            pr = _sub(p, "Project", ProjectId=project_id)
            _sub(pr, "Name", project_name)
            z.writestr("project.bcfp", _xml(p))
            ext = ET.Element("Extensions")
            for tag, item, vals in (("TopicTypes", "TopicType", sorted({t.type for t in topics}) or ["Issue"]),
                                    ("TopicStatuses", "TopicStatus", ["Open"]),
                                    ("Priorities", "Priority", sorted({t.priority for t in topics})),
                                    ("TopicLabels", "TopicLabel", sorted({l for t in topics for l in t.labels})),
                                    ("Users", "User", [author])):
                if vals:
                    g = _sub(ext, tag)
                    for v in vals:
                        _sub(g, item, v)
            z.writestr("extensions.xml", _xml(ext))
        else:
            p = ET.Element("ProjectExtension")
            pr = _sub(p, "Project", ProjectId=project_id)
            _sub(pr, "Name", project_name)
            _sub(p, "ExtensionSchema", "")
            z.writestr("project.bcfp", _xml(p))
        for t in topics:
            tg = _guid("topic", seed, t.key)
            vg = _guid("viewpoint", seed, t.key) if t.components else None
            z.writestr(f"{tg}/markup.bcf", _markup(t, tg, vg, version, author, now, file_name))
            if vg:
                z.writestr(f"{tg}/viewpoint.bcfv", _viewpoint(t, vg, version, camera))
                n_vp += 1
    return {"path": path, "version": version, "topics": len(topics), "viewpoints": n_vp,
            "components": sum(len(t.components) for t in topics)}


def export_report(report: Dict[str, Any], path: str, version: str = "2.1", author: str = "ddc-tools",
                  per_element: bool = False, max_components: int = 1000,
                  only: Sequence[str] = ("added", "removed", "modified"), camera: Optional[Dict[str, Any]] = None,
                  project_name: Optional[str] = None, file_name: Optional[str] = None) -> Dict[str, Any]:
    topics, kind = topics_from_report(report, per_element, max_components, only)
    if kind == "diff":
        fn = file_name or Path(str((report.get("b") or {}).get("source", ""))).name or None
        name = project_name or "ddc diff"
        seed = f"diff|{report.get('a', {}).get('source')}|{report.get('b', {}).get('source')}"
    else:
        fn = file_name or (report.get("source_files") or [None])[0] or Path(str(report.get("source", ""))).name or None
        name = project_name or f"ddc check: {report.get('ruleset') or ''}".strip()
        seed = f"check|{report.get('source')}|{report.get('ruleset')}"
    res = write_bcf(topics, path, version, author, name, fn, camera, seed)
    res["kind"] = kind
    return res


def load_report(path: str) -> Dict[str, Any]:
    try:
        with open(path, encoding="utf-8-sig") as f:
            return json.load(f)
    except (OSError, ValueError) as e:
        raise BcfError(f"cannot read report {path}: {e}")


# ------------------------------------------------------------------------------------------ reader

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _parse(data: bytes, name: str) -> ET.Element:
    head = data[:4096].lower()
    if b"<!doctype" in head or b"<!entity" in data.lower():
        raise BcfError(f"{name}: DOCTYPE / entity declarations are not allowed in BCF files")
    try:
        return ET.fromstring(data)
    except ET.ParseError as e:
        raise BcfError(f"{name}: invalid XML ({e})")


def _kid(e: ET.Element, name: str) -> Optional[ET.Element]:
    for c in e:
        if _local(c.tag) == name:
            return c
    return None


def _kids(e: ET.Element, name: str) -> List[ET.Element]:
    return [c for c in e if _local(c.tag) == name]


def _text(e: Optional[ET.Element], name: str) -> str:
    c = _kid(e, name) if e is not None else None
    return (c.text or "").strip() if c is not None else ""


def _components(e: Optional[ET.Element]) -> List[Dict[str, str]]:
    out = []
    if e is None:
        return out
    for c in e.iter():
        if _local(c.tag) == "Component":
            out.append({"IfcGuid": c.get("IfcGuid") or c.get("ifcGuid") or "", "AuthoringToolId": _text(c, "AuthoringToolId"),
                        "OriginatingSystem": _text(c, "OriginatingSystem")})
    return out


def _viewpoint_info(root: ET.Element) -> Dict[str, Any]:
    comps = _kid(root, "Components")
    info: Dict[str, Any] = {"camera": "", "selection": [], "coloring": [], "default_visibility": None, "visible_exceptions": []}
    for cam in ("PerspectiveCamera", "OrthogonalCamera"):
        if _kid(root, cam) is not None:
            info["camera"] = cam.replace("Camera", "").lower()
    if comps is not None:
        info["selection"] = _components(_kid(comps, "Selection"))
        vis = _kid(comps, "Visibility")
        if vis is not None:
            info["default_visibility"] = (vis.get("DefaultVisibility") or "").lower() or None
            info["visible_exceptions"] = _components(_kid(vis, "Exceptions"))
        col = _kid(comps, "Coloring")
        if col is not None:
            for c in _kids(col, "Color"):
                info["coloring"].append({"color": c.get("Color") or "", "components": _components(c)})
    return info


def read_bcf(path: str) -> Dict[str, Any]:
    try:
        z = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as e:
        raise BcfError(f"{path}: not a BCF zip file ({e})")
    with z:
        infos = z.infolist()
        if len(infos) > MAX_ZIP_MEMBERS:
            raise BcfError("too many files in the archive")
        for i in infos:
            if i.file_size > MAX_ZIP_MEMBER:
                raise BcfError(f"{i.filename}: member too large")
        names = {i.filename: i for i in infos if not i.is_dir()}

        def rd(n: str) -> bytes:
            return z.read(names[n])

        version = ""
        if "bcf.version" in names:
            v = _parse(rd("bcf.version"), "bcf.version")
            version = v.get("VersionId") or _text(v, "DetailedVersion")
        project = {"id": "", "name": ""}
        if "project.bcfp" in names:
            p = _parse(rd("project.bcfp"), "project.bcfp")
            for e in p.iter():
                if _local(e.tag) == "Project":
                    project = {"id": e.get("ProjectId") or "", "name": _text(e, "Name")}
        topics = []
        for n in sorted(names):
            if not n.endswith("markup.bcf") or n.count("/") > 1:
                continue
            folder = n[: -len("markup.bcf")]
            root = _parse(rd(n), n)
            top = _kid(root, "Topic")
            if top is None:
                continue
            hdr = _kid(root, "Header")
            files = []
            if hdr is not None:
                for f in hdr.iter():
                    if _local(f.tag) == "File":
                        files.append(_text(f, "Filename") or _text(f, "Reference"))
            comments = []
            for c in root.iter():
                if _local(c.tag) == "Comment" and c.get("Guid"):
                    vp = _kid(c, "Viewpoint")
                    comments.append({"guid": c.get("Guid"), "date": _text(c, "Date"), "author": _text(c, "Author"),
                                     "text": _text(c, "Comment"), "viewpoint": vp.get("Guid") if vp is not None else "",
                                     "modified": _text(c, "ModifiedDate")})
            vps = []
            for v in root.iter():
                if _local(v.tag) in ("Viewpoints", "ViewPoint") and v.get("Guid"):
                    f = _text(v, "Viewpoint")
                    entry: Dict[str, Any] = {"guid": v.get("Guid"), "file": f, "snapshot": _text(v, "Snapshot")}
                    full = folder + f if f else ""
                    if full in names:
                        try:
                            entry.update(_viewpoint_info(_parse(rd(full), full)))
                        except BcfError as e:
                            entry["error"] = str(e)
                    elif f:
                        entry["error"] = "viewpoint file missing"
                    vps.append(entry)
            labels = [(l.text or "").strip() for l in top.iter() if _local(l.tag) in ("Labels", "Label") and (l.text or "").strip()]
            related = [r.get("Guid") for r in top.iter() if _local(r.tag) == "RelatedTopic" and r.get("Guid")]
            topics.append({
                "guid": top.get("Guid") or folder.strip("/"), "title": _text(top, "Title"),
                "type": top.get("TopicType") or "", "status": top.get("TopicStatus") or "",
                "priority": _text(top, "Priority"), "labels": labels, "stage": _text(top, "Stage"),
                "created": _text(top, "CreationDate"), "author": _text(top, "CreationAuthor"),
                "modified": _text(top, "ModifiedDate"), "modified_by": _text(top, "ModifiedAuthor"),
                "due": _text(top, "DueDate"), "assigned_to": _text(top, "AssignedTo"),
                "description": _text(top, "Description"), "index": _text(top, "Index"),
                "files": [f for f in files if f], "related": related, "comments": comments, "viewpoints": vps,
            })
    return {"tool": "ddc bcf read", "source": str(path), "version": version, "project": project, "topics": topics,
            "summary": _summary(topics)}


def _count(vals: Iterable[str]) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for v in vals:
        out[v or "(none)"] = out.get(v or "(none)", 0) + 1
    return dict(sorted(out.items()))


def _summary(topics: List[Dict[str, Any]]) -> Dict[str, Any]:
    comps = set()
    for t in topics:
        for v in t["viewpoints"]:
            for c in v.get("selection", []):
                comps.add(c["IfcGuid"] or c["AuthoringToolId"])
    return {"topics": len(topics), "comments": sum(len(t["comments"]) for t in topics),
            "viewpoints": sum(len(t["viewpoints"]) for t in topics), "components": len(comps - {""}),
            "by_status": _count(t["status"] for t in topics), "by_priority": _count(t["priority"] for t in topics),
            "by_type": _count(t["type"] for t in topics)}


def write_read_xlsx(res: Dict[str, Any], path: str) -> None:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font

    def x(v):
        if isinstance(v, str):
            return ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]
        return v

    s = res["summary"]
    wb, ws = new_workbook("Summary")
    title_cell(ws, "A1", "BCF summary")
    rows = [("File", Path(res["source"]).name), ("BCF version", res["version"] or "(not stated)"),
            ("Project", res["project"]["name"] or res["project"]["id"]), ("Topics", s["topics"]),
            ("Comments", s["comments"]), ("Viewpoints", s["viewpoints"]), ("Distinct components", s["components"])]
    r = 3
    for k, v in rows:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1
    for title, d in (("By status", s["by_status"]), ("By priority", s["by_priority"]), ("By type", s["by_type"])):
        r += 1
        ws.cell(r, 1, title).font = Font(bold=True)
        for k, v in d.items():
            r += 1
            ws.cell(r, 1, k)
            ws.cell(r, 2, v)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 50

    wt = wb.create_sheet("Topics")
    head = ["Title", "Status", "Priority", "Type", "Labels", "Created", "Author", "Assigned to", "Due", "Stage",
            "Description", "Comments", "Viewpoints", "Components", "Models", "Topic Guid"]
    trs = []
    for t in res["topics"]:
        ncomp = len({c["IfcGuid"] or c["AuthoringToolId"] for v in t["viewpoints"] for c in v.get("selection", [])} - {""})
        trs.append([x(t["title"]), t["status"], t["priority"], t["type"], ", ".join(t["labels"]), t["created"], t["author"],
                    t["assigned_to"], t["due"], t["stage"], x(t["description"]), len(t["comments"]), len(t["viewpoints"]),
                    ncomp, "; ".join(t["files"]), t["guid"]])
    write_table(wt, head, trs)
    autosize(wt)
    wt.column_dimensions["K"].width = 70

    wc = wb.create_sheet("Comments")
    crs = [[x(t["title"]), c["date"], c["author"], x(c["text"]), c["viewpoint"], c["modified"]]
           for t in res["topics"] for c in t["comments"]]
    write_table(wc, ["Topic", "Date", "Author", "Comment", "Viewpoint Guid", "Modified"], crs)
    autosize(wc)

    wk = wb.create_sheet("Components")
    krs = []
    for t in res["topics"]:
        for v in t["viewpoints"]:
            for c in v.get("selection", []):
                krs.append([x(t["title"]), v["guid"], "selected", c["IfcGuid"], c["AuthoringToolId"], c["OriginatingSystem"], ""])
            for col in v.get("coloring", []):
                for c in col["components"]:
                    krs.append([x(t["title"]), v["guid"], "colored", c["IfcGuid"], c["AuthoringToolId"], c["OriginatingSystem"], col["color"]])
            for c in v.get("visible_exceptions", []):
                krs.append([x(t["title"]), v["guid"], "visibility exception", c["IfcGuid"], c["AuthoringToolId"], c["OriginatingSystem"], ""])
    write_table(wk, ["Topic", "Viewpoint Guid", "Role", "IfcGuid", "AuthoringToolId", "OriginatingSystem", "Color"], krs)
    autosize(wk)
    wb.save(path)
