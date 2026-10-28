"""Write diff outputs: xlsx, HTML summary, colored GLB, viewer page."""
from __future__ import annotations
import html
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill

from .diff import DiffResult
from .glb import write_colored_glb
from .model import Model, eid
from .viewer import viewer_html

GREEN, RED, YELLOW, GREY = (0.17, 0.63, 0.17, 1.0), (0.84, 0.15, 0.16, 1.0), (0.95, 0.77, 0.0, 1.0), (0.7, 0.7, 0.7, 0.18)


def _lab(r):
    return [r.get("ElementId"), r.get("UniqueId"), r.get("Category"), r.get("Family"), r.get("Type Name"), r.get("Level")]


LAB = ["ElementId", "UniqueId", "Category", "Family", "Type Name", "Level"]


def write_xlsx(d: DiffResult, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Count"])
    for k, v in d.summary().items():
        ws.append([k, v])
    ws.append(["matched by ElementId only", d.matched_by_id])
    sheets = [("Added", LAB, [_lab(r) for r in d.added]), ("Removed", LAB, [_lab(r) for r in d.removed])]
    for n, h, rows in sheets:
        w = wb.create_sheet(n)
        w.append(h)
        for r in rows:
            w.append(r)
    w = wb.create_sheet("Changed")
    w.append(LAB + ["Parameter changes", "Geometry"])
    for c in d.changed:
        g = c.geom
        gs = "" if not g else ", ".join(x for x in ("moved" if g.moved else "", "resized" if g.resized else "", g.note,
                                                    "triangles %d->%d" % (g.tris_old, g.tris_new) if g.tris_old != g.tris_new else "") if x)
        w.append(_lab(c.new) + [len(c.params), gs])
    w = wb.create_sheet("Parameters")
    w.append(LAB + ["Parameter", "Old", "New"])
    for c in d.changed:
        for p, o, n in c.params:
            w.append(_lab(c.new) + [p, o, n])
    w = wb.create_sheet("Geometry")
    w.append(LAB + ["Moved", "Resized", "Shift x", "Shift y", "Shift z", "Size old", "Size new",
                    "Volume old", "Volume new", "Volume delta", "Triangles old", "Triangles new", "Note"])
    for c in d.changed:
        g = c.geom
        if g:
            w.append(_lab(c.new) + [g.moved, g.resized, *g.shift, str(g.size_old), str(g.size_new),
                                    round(g.volume_old, 6), round(g.volume_new, 6), round(g.volume_delta, 6),
                                    g.tris_old, g.tris_new, g.note])
    for s in wb:
        for c in s[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        s.freeze_panes = "A2"
    wb.save(path)


def _tbl(head, rows, cap=500):
    h = "".join("<th>%s</th>" % html.escape(str(x)) for x in head)
    b = "".join("<tr>%s</tr>" % "".join("<td>%s</td>" % html.escape("" if v is None else str(v)) for v in r)
                for r in rows[:cap])
    more = "<p class=m>%d more rows in the xlsx report</p>" % (len(rows) - cap) if len(rows) > cap else ""
    return "<table><tr>%s</tr>%s</table>%s" % (h, b, more)


CSS = """:root{--bg:#fff;--fg:#1d1d1f;--mut:#666;--ln:#ddd;--hd:#f0f0f0}
@media(prefers-color-scheme:dark){:root{--bg:#1b1b1d;--fg:#e8e8e8;--mut:#999;--ln:#3a3a3d;--hd:#2a2a2d}}
body{background:var(--bg);color:var(--fg);font:14px system-ui,sans-serif;margin:0 auto;max-width:1100px;padding:16px}
table{border-collapse:collapse;width:100%;margin:8px 0 24px}th,td{border:1px solid var(--ln);padding:4px 8px;text-align:left;font-size:13px}
th{background:var(--hd)}.m{color:var(--mut)}.k{display:inline-block;padding:6px 14px;margin:0 8px 8px 0;border:1px solid var(--ln);border-radius:6px}
@media(max-width:600px){table{display:block;overflow-x:auto}}"""


def write_html(d: DiffResult, path, old_name, new_name, viewer=None):
    s = d.summary()
    kp = "".join('<span class=k><b>%d</b> %s</span>' % (v, k) for k, v in s.items())
    pr = [_lab(c.new)[:1] + _lab(c.new)[2:5] + [p, o, n] for c in d.changed for p, o, n in c.params]
    gr = [_lab(c.new)[:1] + _lab(c.new)[2:5] + ["moved" if c.geom.moved else "", "resized" if c.geom.resized else "",
          "%s" % (c.geom.shift,), round(c.geom.volume_delta, 5), "%d->%d" % (c.geom.tris_old, c.geom.tris_new), c.geom.note]
          for c in d.changed if c.geom]
    lab = ["ElementId", "Category", "Family", "Type Name"]
    body = ["<h1>Model diff</h1><p class=m>%s &rarr; %s</p>%s" % (html.escape(old_name), html.escape(new_name), kp)]
    if viewer:
        body.append('<p><a href="%s">Open 3D viewer</a> (green added, red removed, yellow changed)</p>' % html.escape(viewer))
    body += ["<h2>Added</h2>" + _tbl(lab, [_lab(r)[:1] + _lab(r)[2:5] for r in d.added]),
             "<h2>Removed</h2>" + _tbl(lab, [_lab(r)[:1] + _lab(r)[2:5] for r in d.removed]),
             "<h2>Parameter changes</h2>" + _tbl(lab + ["Parameter", "Old", "New"], pr),
             "<h2>Geometry changes</h2>" + _tbl(lab + ["Moved", "Resized", "Centre shift (m)", "Volume delta", "Triangles", "Note"], gr)]
    open(path, "w", encoding="utf-8").write(
        '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        "<title>Model diff</title><style>%s</style></head><body>%s</body></html>" % (CSS, "".join(body)))


def write_glb(d: DiffResult, old: Model, new: Model, path, context=True):
    items = []
    for r in d.added:
        m = new.meshes.get(eid(r))
        if m:
            items.append(("added:%s" % eid(r), m, GREEN))
    for r in d.removed:
        m = old.meshes.get(eid(r))
        if m:
            items.append(("removed:%s" % eid(r), m, RED))
    ch = set()
    for c in d.changed:
        m = new.meshes.get(eid(c.new)) or old.meshes.get(eid(c.old))
        ch.add(eid(c.new))
        if m:
            items.append(("changed:%s" % eid(c.new), m, YELLOW))
    if context:
        ad = {eid(r) for r in d.added}
        for e, m in new.meshes.items():
            if e not in ch and e not in ad:
                items.append(("unchanged:%s" % e, m, GREY))
    write_colored_glb(path, items)
    return {k: sum(1 for n, _, _ in items if n.startswith(k + ":")) for k in ("added", "removed", "changed", "unchanged")}


def write_all(d, old, new, outdir, old_name, new_name):
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    write_xlsx(d, out / "diff.xlsx")
    counts = write_glb(d, old, new, out / "diff.glb")
    viewer_html(out / "diff.glb", "Model diff", counts, out / "viewer.html")
    write_html(d, out / "diff.html", old_name, new_name, "viewer.html")
    return out
