"""Load a model: RVT (converted via the DDC CLI), a *.converted folder, or elements.xlsx."""
from __future__ import annotations
import hashlib
import os
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
import openpyxl

from .glb import element_meshes

ID_COLS = ("ElementId", "GlobalId", "UniqueId")
DEFAULT_CLI = r"C:\tmp\r205\pkg\General\cli\RVTconverter.exe"


@dataclass
class Model:
    path: str
    headers: list
    rows: list                                   # list[dict]
    meshes: dict = field(default_factory=dict)   # ElementId(int) -> Mesh
    glb: str | None = None


def cli_path():
    p = os.environ.get("DDC_RVT_CLI") or shutil.which("RVTconverter") or DEFAULT_CLI
    if not os.path.exists(p):
        raise FileNotFoundError("RVT converter not found; set DDC_RVT_CLI")
    return p


def convert(rvt, cache=None, cli=None) -> Path:
    rvt = Path(rvt).resolve()
    st = rvt.stat()
    tag = hashlib.sha1(("%s|%d|%d" % (rvt, st.st_size, st.st_mtime_ns)).encode()).hexdigest()[:10]
    base = Path(cache or os.environ.get("DDC_QA_CACHE") or Path(os.environ.get("TEMP", ".")) / "ddc-qa-cache")
    out = base / ("%s-%s" % (rvt.stem, tag))
    if (out / "elements.xlsx").exists() and (out / "native_graph.glb").exists():
        return out
    out.mkdir(parents=True, exist_ok=True)
    p = subprocess.run([cli or cli_path(), str(rvt), "--xlsx", "--out", str(out)], capture_output=True, text=True)
    if not (out / "elements.xlsx").exists():
        raise RuntimeError("conversion failed (rc=%s): %s" % (p.returncode, (p.stdout + p.stderr)[-400:]))
    return out


def read_elements(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Elements"] if "Elements" in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else "" for h in next(it)]
    rows = []
    for r in it:
        if all(v is None for v in r):
            continue
        rows.append({h: v for h, v in zip(hdr, r) if h})
    wb.close()
    return hdr, rows


def load_model(path, cache=None, cli=None, geometry=True) -> Model:
    p = Path(path)
    if p.suffix.lower() in (".rvt", ".rfa"):
        d = convert(p, cache, cli)
        x = d / "elements.xlsx"
    elif p.is_dir():
        d, x = p, p / "elements.xlsx"
    else:
        d, x = p.parent, p
    hdr, rows = read_elements(x)
    glb = d / "native_graph.glb"
    if not glb.exists():
        g = sorted(d.glob("*.glb"))
        glb = g[0] if g else None
    meshes = element_meshes(glb) if (geometry and glb) else {}
    return Model(str(path), hdr, rows, meshes, str(glb) if glb else None)


def eid(r):
    try:
        return int(r.get("ElementId"))
    except (TypeError, ValueError):
        return None
