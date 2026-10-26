"""Rule-check outputs: xlsx and HTML."""
from __future__ import annotations
import html
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill

from .diffreport import CSS, _tbl
from .rules import Finding


def _row(f: Finding):
    e, o = f.element or {}, f.other or {}
    return [f.rule_id, f.title, f.severity, e.get("ElementId"), e.get("Category"), e.get("Type Name"), e.get("Level"),
            Finding.guid_of(e), f.message, o.get("ElementId"), o.get("Category")]


HEAD = ["Rule", "Title", "Severity", "ElementId", "Category", "Type Name", "Level", "IfcGUID", "Message",
        "Other ElementId", "Other Category"]


def write_xlsx(findings, path, notes=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Rule", "Title", "Severity", "Findings"])
    cnt = Counter((f.rule_id, f.title, f.severity) for f in findings)
    for (r, t, s), n in sorted(cnt.items()):
        ws.append([r, t, s, n])
    for n in notes:
        ws.append(["note", n])
    w = wb.create_sheet("Findings")
    w.append(HEAD)
    for f in findings:
        w.append(_row(f))
    for s in wb:
        for c in s[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        s.freeze_panes = "A2"
    wb.save(path)


def write_html(findings, path, title, notes=()):
    sev = Counter(f.severity for f in findings)
    kp = "".join('<span class=k><b>%d</b> %s</span>' % (v, k) for k, v in sev.items()) or "<p>No findings.</p>"
    cnt = Counter((f.rule_id, f.title, f.severity) for f in findings)
    body = ["<h1>Rule check</h1><p class=m>%s</p>%s" % (html.escape(title), kp),
            "<h2>By rule</h2>" + _tbl(["Rule", "Title", "Severity", "Findings"], [list(k) + [v] for k, v in sorted(cnt.items())]),
            "<h2>Findings</h2>" + _tbl(HEAD[:9], [_row(f)[:9] for f in findings], cap=1000)]
    if notes:
        body.append("<h2>Notes</h2><ul>%s</ul>" % "".join("<li>%s</li>" % html.escape(n) for n in notes))
    open(path, "w", encoding="utf-8").write(
        '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        "<title>Rule check</title><style>%s</style></head><body>%s</body></html>" % (CSS, "".join(body)))
