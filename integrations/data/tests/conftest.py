import openpyxl
import pytest

HEAD = ["ElementId", "GlobalId", "UniqueId", "Category", "IfcClass", "Family", "Type Name", "Level",
        "Length (m)", "Width (m)", "Height (m)", "Area (m²)", "Volume (m³)", "Fire Rating",
        "Door Material", "Structural Material", "Mark", "Thickness (internal units)"]


def make_xlsx(path, header=HEAD, rows=(), base=True, extra_sheets=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elements"
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    if base:
        b = wb.create_sheet("_ddc_base")
        b.append(["format", "ddc-excel-base/1"])
        b.append(["model_sha256", "abc"])
        b.append(["data_sheet", "Elements"])
        b.append(header)
        for r in rows:
            b.append(list(r))
        b.sheet_state = "hidden"
    for name, (h, rr) in (extra_sheets or {}).items():
        s = wb.create_sheet(name)
        s.append(h)
        for r in rr:
            s.append(r)
    wb.save(path)
    return str(path)


def row(i, cat, fam="F", typ="T", lvl="L1", length=None, w=None, h=None, area=None, vol=None, fire=None,
        dmat=None, smat=None):
    return [i, f"G{i}", f"U{i}", cat, None, fam, typ, lvl, length, w, h, area, vol, fire, dmat, smat, None, None]


ROWS = [
    row(1, "Walls", "Basic Wall", "Generic - 200mm", "L1", 10.0, 0.2, 3.0, 30.0, 6.0, "1h", None, "Concrete"),
    row(2, "Walls", "Basic Wall", "Generic - 200mm", "L2", 5.0, 0.2, 3.0, 15.0, 3.0, None, None, "Concrete"),
    row(3, "Walls", "Basic Wall", "Brick - 300mm", "L1", 4.0, 0.3, 3.0, 12.0, 3.6, None, None, "Brick"),
    row(4, "Doors", "Single", "0900 x 2100", "L1", None, 0.9, 2.1, 1.89, None, None, "Oak", None),
    row(5, "Doors", "Single", "0900 x 2100", "L2", None, 0.9, 2.1, 1.89, None, "30m", "Oak", None),
    row(6, "Windows", "Fixed", "1200 x 1200", "L1", None, 1.2, 1.2, 1.44, None, None, None, None),
    row(7, "Rooms", "Room", "Büro Ä/Ö", "L1", None, None, None, 25.5, 76.5),
    row(8, "Furniture", None, None, None),
]


@pytest.fixture
def xlsx(tmp_path):
    return make_xlsx(tmp_path / "m.xlsx", rows=ROWS)
