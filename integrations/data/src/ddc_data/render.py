"""Render a Report to formatted xlsx (openpyxl) and self-contained HTML."""
from __future__ import annotations

import html
import math
import re

import pandas as pd

from .report import Report, Section


def _label(sec: Section, col: str) -> str:
    u = sec.units.get(col)
    return f"{col} ({u})" if u else str(col)


def _clean(v):
    if v is None or v is pd.NA or v is pd.NaT or (isinstance(v, float) and math.isnan(v)):
        return None
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        v = v.item()
    return v


def _sheet_name(title: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", title)[:31] or "Sheet"
    n, k = base, 1
    while n.lower() in used:
        k += 1
        n = f"{base[:28]}_{k}"
    used.add(n.lower())
    return n


def render_xlsx(rep: Report, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Summary"
    ws0["A1"] = rep.name
    ws0["A1"].font = Font(bold=True, size=14)
    ws0["A2"] = rep.description
    ws0["A3"] = f"Source: {rep.source}"
    ws0["A5"], ws0["B5"] = "Section", "Rows"
    for c in ("A5", "B5"):
        ws0[c].font = Font(bold=True)
    used = {"summary"}
    head_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="BFBFBF")
    for i, sec in enumerate(rep.sections):
        ws = wb.create_sheet(_sheet_name(sec.title, used))
        ws0.cell(row=6 + i, column=1, value=sec.title)
        ws0.cell(row=6 + i, column=2, value=len(sec.df))
        cols = list(sec.df.columns)
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j, value=_label(sec, c))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths = [len(_label(sec, c)) for c in cols]
        for r, row in enumerate(sec.df.itertuples(index=False), 2):
            for j, (c, v) in enumerate(zip(cols, row), 1):
                v = _clean(v)
                cell = ws.cell(row=r, column=j, value=v)
                if isinstance(v, float):
                    cell.number_format = "0.0%" if sec.formats.get(c) == "pct" else "#,##0.00"
                elif isinstance(v, int) and not isinstance(v, bool):
                    cell.number_format = "#,##0" if sec.formats.get(c) == "int" else "0"
                if v is not None:
                    widths[j - 1] = max(widths[j - 1], min(len(str(v)), 60))
        last = len(sec.df) + 1
        if sec.totals and len(sec.df):
            tr = last + 1
            ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
            for j, c in enumerate(cols, 1):
                if c in sec.totals:
                    cell = ws.cell(row=tr, column=j, value=sec.totals[c])
                    cell.font = Font(bold=True)
                    cell.number_format = "#,##0" if sec.formats.get(c) == "int" else "#,##0.00"
                    cell.border = Border(top=thin)
        if cols:
            ws.freeze_panes = "A2"
            if len(sec.df):
                ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(8, min(w + 2, 62))
        if not len(sec.df):
            ws.cell(row=2, column=1, value="(no rows)")
        if sec.note:
            ws.cell(row=last + 3, column=1, value=sec.note)
    ws0.column_dimensions["A"].width = 40
    wb.save(path)


_CSS = ("body{font:14px/1.4 system-ui,Segoe UI,sans-serif;margin:24px;color:#1b1b1b;background:#fff}"
        "h1{margin:0 0 4px}h2{margin:28px 0 8px}p.d{color:#555;margin:0 0 8px}"
        "table{border-collapse:collapse;min-width:50%}th{background:#1f4e78;color:#fff;padding:6px 10px;text-align:left}"
        "td{padding:4px 10px;border-bottom:1px solid #ddd}td.n{text-align:right;font-variant-numeric:tabular-nums}"
        "tr.t td{font-weight:700;border-top:2px solid #888}.note{color:#8a5a00}"
        "@media (prefers-color-scheme:dark){body{background:#161616;color:#e6e6e6}td{border-color:#333}"
        "th{background:#2d5f8f}p.d{color:#aaa}}")


def _fmt(v, kind):
    v = _clean(v)
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v * 100:.1f}%" if kind == "pct" else f"{v:,.2f}"
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}" if kind == "int" else str(v)
    return str(v)


def render_html(rep: Report, path: str) -> None:
    e = html.escape
    out = ["<!doctype html><html lang=en><meta charset=utf-8>"
           "<meta name=viewport content='width=device-width,initial-scale=1'>"
           f"<title>{e(rep.name)}</title><style>{_CSS}</style><body>",
           f"<h1>{e(rep.name)}</h1><p class=d>{e(rep.description)}</p><p class=d>Source: {e(rep.source)}</p>"]
    for sec in rep.sections:
        out.append(f"<h2>{e(sec.title)} <small>({len(sec.df)} rows)</small></h2>")
        if sec.note:
            out.append(f"<p class=note>{e(sec.note)}</p>")
        if not len(sec.df):
            out.append("<p>(no rows)</p>")
            continue
        cols = list(sec.df.columns)
        out.append("<div style='overflow-x:auto'><table><tr>" +
                   "".join(f"<th>{e(_label(sec, c))}</th>" for c in cols) + "</tr>")
        for row in sec.df.itertuples(index=False):
            out.append("<tr>" + "".join(
                f"<td class={'n' if isinstance(_clean(v), (int, float)) else 's'}>{e(_fmt(v, sec.formats.get(c)))}</td>"
                for c, v in zip(cols, row)) + "</tr>")
        if sec.totals:
            out.append("<tr class=t>" + "".join(
                "<td class=n>" + e(_fmt(sec.totals[c], sec.formats.get(c)) if c in sec.totals
                                   else ("Total" if i == 0 else "")) + "</td>"
                for i, c in enumerate(cols)) + "</tr>")
        out.append("</table></div>")
    out.append("</body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("".join(out))
