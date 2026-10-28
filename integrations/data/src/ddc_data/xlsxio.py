"""Read the DDC xlsx output into typed Arrow tables (one per visible sheet)."""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import pyarrow as pa

ID_COLUMNS = {"ElementId", "Element Id", "Id", "TypeId", "LevelId", "OwnerViewId"}
_UNIT_RE = re.compile(r"^(?P<base>.*?)\s*\((?P<unit>m|m²|m³|mm|cm|ft|ft²|ft³|in|kg|t|°|deg|%|internal units)\)\s*$")
META_SHEET = "_ddc_base"


def split_unit(header: str) -> tuple[str, str | None]:
    """'Area (m²)' -> ('Area', 'm²'); headers without a unit suffix pass through."""
    m = _UNIT_RE.match(header)
    return (m.group("base"), m.group("unit")) if m else (header, None)


@dataclass
class Workbook:
    tables: dict[str, pa.Table] = field(default_factory=dict)
    meta: dict[str, str] = field(default_factory=dict)


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _column(name: str, vals: list) -> pa.Array:
    nn = [v for v in vals if v is not None]
    if not nn:
        return pa.nulls(len(vals), pa.string())
    if all(_is_num(v) for v in nn):
        if name in ID_COLUMNS and all(float(v).is_integer() for v in nn):
            return pa.array([None if v is None else int(v) for v in vals], pa.int64())
        return pa.array([None if v is None else float(v) for v in vals], pa.float64())
    if all(isinstance(v, bool) for v in nn):
        return pa.array(vals, pa.bool_())
    return pa.array([None if v is None else (v if isinstance(v, str) else str(v)) for v in vals], pa.string())


def _dedupe(headers: list) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(headers):
        h = str(h) if h not in (None, "") else f"column_{i + 1}"
        n = seen.get(h, 0)
        seen[h] = n + 1
        out.append(h if n == 0 else f"{h}.{n}")
    return out


def build_table(headers: list, rows) -> pa.Table:
    names = _dedupe(headers)
    cols: list[list] = [[] for _ in names]
    w = len(names)
    for r in rows:
        r = tuple(r)
        for i in range(w):
            cols[i].append(r[i] if i < len(r) else None)
    fields, arrays = [], []
    for n, vals in zip(names, cols):
        arr = _column(n, vals)
        base, unit = split_unit(n)
        md = {b"ddc_base_name": base.encode()}
        if unit:
            md[b"unit"] = unit.encode()
        fields.append(pa.field(n, arr.type, metadata=md))
        arrays.append(arr)
    return pa.Table.from_arrays(arrays, schema=pa.schema(fields))


def read_workbook(path: str, include_base: bool = False) -> Workbook:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = Workbook()
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            first = next(it, None)
            if ws.title == META_SHEET:
                # key/value preamble; the rest mirrors Elements and is skipped
                rows = [first] if first else []
                for r in it:
                    if r and r[0] is not None and all(c is None for c in r[2:]) and len(rows) < 64:
                        rows.append(r)
                    else:
                        break
                for r in rows:
                    if r and r[0] is not None and len(r) > 1 and r[1] is not None:
                        out.meta[str(r[0])] = str(r[1])
                if not include_base:
                    continue
                it = ws.iter_rows(values_only=True)
                first = next(it, None)
            elif ws.sheet_state != "visible":
                continue
            if first is None:
                out.tables[ws.title] = pa.table({})
                continue
            out.tables[ws.title] = build_table(list(first), it)
    finally:
        wb.close()
    return out


def read_table(path: str, sheet: str = "Elements") -> pa.Table:
    wb = read_workbook(path)
    if sheet in wb.tables:
        return wb.tables[sheet]
    if wb.tables:
        return next(iter(wb.tables.values()))
    return pa.table({})
